"""Import the edits you made in the workbook, field by field, without ever trusting the file blindly.

Rules:
- Only designated columns are read (status, date applied, next action and its date, interview
  date, notes; on contacts: interaction status, next follow-up date, next action, notes).
  Scores, URLs, IDs and everything else are read-only and ignored.
- Each row carries the record version it was exported with. If the database moved on since
  then and the row also differs, that row is a conflict: nothing is applied for it unless you
  explicitly choose the workbook (`prefer_excel`).
- Formulas, macros and unknown IDs are refused; values are validated before anything changes.
- A database backup is taken before the first change; re-importing the same file is a no-op
  (idempotency comes from comparing values, so a real re-application is always logged).
"""

from __future__ import annotations

import json
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from itertools import pairwise
from pathlib import Path
from typing import Any, Callable
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import load_workbook

from summer_hunter.crm import statuses as st
from summer_hunter.crm.service import CRM, _same
from summer_hunter.crm.validation import (
    MAX_NOTE,
    ValidationError,
    app_code,
    clean_text,
    contact_code,
    excel_unescape,
    parse_app_id,
    parse_contact_id,
    parse_date,
    parse_datetime,
)
from summer_hunter.database.repository import Repository, as_utc

APP_FIELDS = {   # workbook header -> (attribute, parser kind)
    "Application Status": ("status", "status"),
    "Date Applied": ("date_applied", "date"),
    "Next Action": ("next_action", "text"),
    "Next Action Date": ("next_action_date", "date"),
    "Interview Date": ("interview_at", "datetime"),
    "Notes": ("notes", "note"),
}
CONTACT_FIELDS = {
    "Interaction Status": ("interaction_status", "contact_status"),
    "Next Follow-up Date": ("next_followup_date", "date"),
    "Next Action": ("next_action", "text"),
    "Notes": ("notes", "note"),
}


@dataclass
class Conflict:
    record: str
    field: str
    workbook: Any
    database: Any


@dataclass
class ImportReport:
    applied: list[str] = field(default_factory=list)
    conflicts: list[Conflict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    ignored: list[str] = field(default_factory=list)
    backup: Path | None = None
    dry_run: bool = False

    @property
    def changed(self) -> bool:
        return bool(self.applied)

    def summary(self) -> str:
        parts = [f"{len(self.applied)} change(s) applied" if not self.dry_run else f"{len(self.applied)} change(s) pending"]
        if self.conflicts:
            parts.append(f"{len(self.conflicts)} conflict(s)")
        if self.errors:
            parts.append(f"{len(self.errors)} error(s)")
        if self.ignored:
            parts.append(f"{len(self.ignored)} row(s) ignored")
        return ", ".join(parts)


class WorkbookError(ValueError):
    pass


def _open(path: Path):  # noqa: ANN202
    if path.suffix.lower() != ".xlsx":
        raise WorkbookError("only .xlsx workbooks are accepted (no macro-enabled files)")
    try:
        with zipfile.ZipFile(path) as z:
            if any(n.lower().endswith("vbaproject.bin") for n in z.namelist()):
                raise WorkbookError("the workbook contains macros; refused")
    except zipfile.BadZipFile:
        raise WorkbookError("not a valid .xlsx file") from None
    wb = load_workbook(path, data_only=False, keep_vba=False, keep_links=False)
    if "_meta" not in wb.sheetnames or "Applications" not in wb.sheetnames:
        raise WorkbookError("this is not a Summer Hunter CRM workbook (missing _meta or Applications sheet)")
    return wb


def _headers(ws, required: list[str]) -> dict[str, int]:  # noqa: ANN001
    found = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    missing = [h for h in required if h not in found]
    if missing:
        raise WorkbookError(f"sheet {ws.title}: missing column(s) {', '.join(missing)}")
    return found


def _parse(kind: str, raw: Any, header: str, tz: str) -> Any:
    raw = excel_unescape(raw)
    if kind == "status":
        value = str(raw or "").strip().upper()
        if value not in st.STATUSES:
            raise ValidationError(f"{header}: {raw!r} is not a valid status")
        return value
    if kind == "contact_status":
        value = str(raw or "").strip().upper()
        if value not in st.INTERACTION_STATUSES:
            raise ValidationError(f"{header}: {raw!r} is not a valid interaction status")
        return value
    if kind == "date":
        return parse_date(raw, header)
    if kind == "datetime":
        return parse_datetime(raw, tz, header)
    if kind == "note":
        return clean_text(raw, header, MAX_NOTE)
    return clean_text(raw, header)


def _read_rows(ws, id_header: str, fields: dict[str, tuple[str, str]], tz: str, report: ImportReport,  # noqa: ANN001
               parse_id: Callable[[Any], int]) -> list[tuple[int, int | None, dict[str, Any], int]]:
    cols = _headers(ws, [id_header, "Row Version", *fields])
    out = []
    seen: dict[int, list[int]] = {}
    for row in ws.iter_rows(min_row=2):
        cells = {h: row[c - 1] for h, c in cols.items() if c - 1 < len(row)}
        raw_id = cells[id_header].value
        if raw_id in (None, ""):
            if any(cells[h].value not in (None, "") for h in fields):
                report.ignored.append(f"{ws.title} row {row[0].row}: no ID (new rows are not imported; add records "
                                      "through Telegram or the CLI)")
            continue
        try:
            rid = parse_id(raw_id)
            values = {}
            for header, (attr, kind) in fields.items():
                cell = cells[header]
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    raise ValidationError(f"{header}: formulas are not accepted")
                values[attr] = _parse(kind, cell.value, header, tz)
            version = cells["Row Version"].value
            version = int(version) if version not in (None, "") else None
        except (ValidationError, ValueError) as exc:
            report.errors.append(f"{ws.title} row {row[0].row} ({raw_id}): {exc}")
            continue
        seen.setdefault(rid, []).append(row[0].row)
        out.append((rid, version, values, row[0].row))
    duplicated = {rid for rid, lines in seen.items() if len(lines) > 1}
    for rid in sorted(duplicated):     # ambiguous: which row is the truth? none of them is applied
        report.errors.append(f"{ws.title}: ID {raw_label(id_header, rid)} appears on rows {seen[rid]}; fix the duplicate")
    return [r for r in out if r[0] not in duplicated]


def raw_label(id_header: str, rid: int) -> str:
    return app_code(rid) if id_header == "Application ID" else contact_code(rid)


def display(value: Any) -> Any:
    if isinstance(value, datetime):
        return as_utc(value).isoformat(timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    return value


def merge(record: Any, code: str, values: dict[str, Any], version: int | None, base: dict[str, Any] | None,
          prefer_excel: bool, report: ImportReport) -> dict[str, Any]:
    """Three-way merge per field: exported value (base), workbook value, current database value.

    A cell you did not touch is ignored even if the database moved on (stale file, no conflict).
    A cell you changed is applied when the database still holds the exported value, skipped when
    the database already has your value, and reported as a conflict otherwise. Without the export
    snapshot (file from elsewhere), the row version decides: any difference on a newer row conflicts.
    """
    out, conflicts = {}, []
    for key, new in values.items():
        current = getattr(record, key)
        if _same(current, new):
            continue
        if base is not None:
            if display(new) == base.get(key):
                continue                              # untouched in the workbook
            if display(current) == base.get(key):
                out[key] = new                        # only the workbook changed
                continue
        elif version == record.version:
            out[key] = new
            continue
        conflicts.append(Conflict(code, key, display(new), display(current)))
    if conflicts and prefer_excel:
        out.update({c.field: values[c.field] for c in conflicts})
    elif conflicts:
        report.conflicts += conflicts
        return {}                                     # a row with a conflict is left untouched as a whole
    return out


def import_workbook(crm: CRM, path: Path, prefer_excel: bool = False, dry_run: bool = False,
                    backup: Callable[[], Path] | None = None) -> ImportReport:
    """Apply authorized edits from `path`. The caller commits (or rolls back) the session."""
    report = ImportReport(dry_run=dry_run)
    wb = _open(path)
    meta = {str(r[0].value): r[1].value for r in wb["_meta"].iter_rows(min_row=1, max_col=2) if r[0].value}
    raw_snap = Repository(crm.s).get_kv(f"crm_export:{meta.get('export_id')}") if meta.get("export_id") else None
    snap: dict[str, dict[str, Any]] = json.loads(raw_snap) if raw_snap else {}
    tz = str(meta.get("timezone") or crm.tz)       # cells were written in the export's timezone
    try:
        ZoneInfo(tz)
    except (ZoneInfoNotFoundError, ValueError):
        raise WorkbookError(f"_meta timezone {tz!r} is not valid") from None
    backed_up = False

    def ensure_backup() -> None:
        nonlocal backed_up
        if not backed_up and backup is not None and not dry_run:
            report.backup = backup()
            backed_up = True

    for rid, version, values, line in _read_rows(wb["Applications"], "Application ID", APP_FIELDS, tz, report, parse_app_id):
        app = crm.application(rid)
        code = app_code(rid)
        if app is None:
            report.ignored.append(f"{code}: unknown application (row {line})")
            continue
        diffs = merge(app, code, values, version, snap.get(code), prefer_excel, report)
        if not diffs:
            continue
        target = diffs.get("status")
        submitting = app.date_applied is None and diffs.get("date_applied") is not None
        if submitting and app.status in st.PRE_APPLICATION and (target is None or target in st.PRE_APPLICATION):
            target = "APPLIED"         # a submission date entered in Excel is a confirmed submission
        try:
            path_ok = [app.status] + (["APPLIED"] if submitting and target not in (None, "APPLIED") and
                                      app.status in st.PRE_APPLICATION else []) + ([target] if target else [])
            for a, b in pairwise(path_ok):
                st.check_transition(a, b, (diffs.get("date_applied", app.date_applied)) is not None)
        except st.TransitionError as exc:
            report.errors.append(f"{code}: {exc}")
            continue
        if dry_run:
            report.applied += [f"{code} {k}: {display(getattr(app, k))} → {display(v)}" for k, v in diffs.items()]
            continue
        ensure_backup()
        try:
            with crm.s.begin_nested():
                fields = {k: v for k, v in diffs.items() if k != "status"}
                if "date_applied" in fields:     # a submission date must exist before a post-application status
                    crm.update_fields(app, {"date_applied": fields.pop("date_applied")}, "excel")
                if submitting and app.status in st.PRE_APPLICATION:
                    crm.set_status(app, "APPLIED", "excel", notes="date applied entered in the workbook")
                if target and target != app.status:
                    crm.set_status(app, target, "excel")
                if fields:
                    crm.update_fields(app, fields, "excel")
                if app.date_applied is not None:
                    crm.freeze_documents(app)
        except (ValidationError, st.TransitionError) as exc:
            report.errors.append(f"{code}: {exc}")
            continue
        report.applied += [f"{code} {k}: {display(v)}" for k, v in diffs.items()]

    if "Networking" in wb.sheetnames:
        for rid, version, values, line in _read_rows(wb["Networking"], "Contact ID", CONTACT_FIELDS, tz, report,
                                                     parse_contact_id):
            contact = crm.contact(rid)
            code = contact_code(rid)
            if contact is None:
                report.ignored.append(f"{code}: unknown contact (row {line})")
                continue
            diffs = merge(contact, code, values, version, snap.get(code), prefer_excel, report)
            if not diffs:
                continue
            if dry_run:
                report.applied += [f"{code} {k}: {display(getattr(contact, k))} → {display(v)}" for k, v in diffs.items()]
                continue
            ensure_backup()
            changed = [f"{code} {k}: {display(v)}" for k, v in diffs.items()]
            try:
                with crm.s.begin_nested():
                    if "interaction_status" in diffs:
                        crm.set_contact_status(contact, diffs.pop("interaction_status"), "excel")
                    if diffs:
                        crm.update_contact_fields(contact, diffs, "excel")
            except ValidationError as exc:
                report.errors.append(f"{code}: {exc}")
                continue
            report.applied += changed
    return report
