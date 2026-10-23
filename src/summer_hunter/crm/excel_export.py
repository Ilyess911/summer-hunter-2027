"""Build Summer_Hunter_2027_CRM.xlsx from SQLite. The workbook is a mirror, never the source of truth.

Every cell holds a plain value written by Python: no formula carries business logic (conditional
formatting rules are display only), untrusted text is neutralized against formula injection, and
only http(s) links become hyperlinks.
"""

from __future__ import annotations

import json
import shutil
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select

from summer_hunter.crm import statuses as st
from summer_hunter.crm.metrics import Metrics, compute
from summer_hunter.crm.service import CRM
from summer_hunter.crm.validation import app_code, contact_code, excel_safe
from summer_hunter.database.models import KeyValue, Opportunity, SourceRecord
from summer_hunter.database.repository import Repository, as_utc
from summer_hunter.engine.geo import country_name
from summer_hunter.notifications.formatter import safe_url

FILENAME = "Summer_Hunter_2027_CRM.xlsx"
FORMAT_VERSION = "1"
KEEP_SNAPSHOTS = 40

NAVY = "1F3864"
NAVY_EDIT = "2F5597"
GRAY = "F2F2F2"
MID_GRAY = "D9D9D9"
TEXT = "262626"
WHITE = "FFFFFF"
DATE_FMT = "yyyy-mm-dd"
DATETIME_FMT = "yyyy-mm-dd hh:mm"

# (header, width, kind). kind: text | int | date | datetime | url | edit-* (editable on import)
APPLICATION_COLUMNS = [
    ("Application ID", 14, "text"), ("Company", 22, "text"), ("Job Title", 38, "text"), ("Job Family", 20, "text"),
    ("Country", 14, "text"), ("City", 14, "text"), ("Career Fit Score", 11, "int"), ("Eligibility Status", 18, "text"),
    ("Application Status", 20, "edit-status"), ("Date Discovered", 13, "date"), ("Application Deadline", 13, "date"),
    ("Date Applied", 13, "edit-date"), ("Last Update", 17, "datetime"), ("Next Action", 30, "edit-text"),
    ("Next Action Date", 13, "edit-date"), ("Interview Date", 17, "edit-datetime"), ("Recruiter Name", 20, "text"),
    ("Recruiter LinkedIn", 26, "url"), ("Job URL", 26, "url"), ("Application URL", 26, "url"), ("CV Version", 16, "text"),
    ("Cover Letter Version", 16, "text"), ("Salary / Compensation", 18, "text"), ("Internship Start Date", 13, "date"),
    ("Internship End Date", 13, "date"), ("Notes", 44, "edit-note"),
    ("Career Category", 22, "text"), ("Recruitability Score", 12, "int"), ("Industrial Priority Score", 12, "int"),
    ("Relevant Industrial Experience", 40, "text"), ("Missing Technical Requirements", 34, "text"),
    ("Compensation Status", 20, "text"), ("Salary Minimum", 12, "num"), ("Salary Maximum", 12, "num"),
    ("Salary Currency", 9, "text"), ("Salary Period", 9, "text"), ("Salary Basis", 9, "text"), ("Monthly Salary EUR", 12, "num"),
    ("Estimated Monthly Living Costs", 12, "num"), ("Estimated Disposable Monthly Balance", 13, "num"),
    ("Housing Benefits", 22, "text"), ("Relocation Benefits", 22, "text"), ("Visa Support", 22, "text"),
    ("Compensation Confidence", 11, "text"), ("Compensation Attractiveness Score", 12, "int"),
    ("Global Application Priority Score", 12, "num"), ("Recommended Next Action", 20, "text"),
    ("Application Strategy Category", 12, "text"), ("Opportunity Source", 18, "text"), ("Last Verification Date", 13, "date"),
    ("Row Version", 8, "int"),
]
CONTACT_COLUMNS = [
    ("Contact ID", 11, "text"), ("Company", 22, "text"), ("Contact Name", 22, "text"), ("Position", 26, "text"),
    ("LinkedIn Profile", 28, "url"), ("Related Application ID", 14, "text"), ("Contact Type", 16, "text"),
    ("First Contact Date", 13, "date"), ("Last Interaction Date", 13, "date"), ("Interaction Status", 20, "edit-status"),
    ("Next Follow-up Date", 13, "edit-date"), ("Next Action", 28, "edit-text"), ("Notes", 40, "edit-note"),
    ("Row Version", 8, "int"),
]
OPPORTUNITY_COLUMNS = [
    ("Opportunity ID", 11, "int"), ("Company", 22, "text"), ("Job Title", 38, "text"), ("Country", 14, "text"),
    ("Career Fit Score", 11, "int"), ("Networking Score", 11, "int"), ("Eligibility Status", 18, "text"),
    ("Date Discovered", 13, "date"), ("Application Deadline", 13, "date"), ("Official Job URL", 28, "url"),
    ("LinkedIn Job URL", 28, "url"), ("LinkedIn Post URL", 28, "url"), ("Opportunity Status", 18, "text"),
    ("Reason for Recommendation", 60, "text"), ("Career Category", 22, "text"), ("Recruitability Score", 12, "int"),
    ("Industrial Priority Score", 12, "int"),
]
EVENT_COLUMNS = [
    ("Timestamp", 17, "datetime"), ("Application ID", 14, "text"), ("Company", 22, "text"), ("Event Type", 22, "text"),
    ("Previous Value", 22, "text"), ("New Value", 22, "text"), ("Source", 11, "text"), ("Notes", 50, "text"),
]

STATUS_FILL = {
    "DISCOVERED": "D9D9D9", "SHORTLISTED": "BDD7EE", "TO_APPLY": "FFF2CC", "PREPARING": "F8CBAD",
    "APPLIED": "DDEBF7", "ONLINE_ASSESSMENT": "E4DFEC", "INTERVIEW": "D9C3E9", "FINAL_INTERVIEW": "C9A9E0",
    "OFFER_RECEIVED": "C6EFCE", "OFFER_ACCEPTED": "A9D08E", "REJECTED": "FFC7CE", "WITHDRAWN": "A6A6A6",
    "CLOSED": "A6A6A6",
}
TERMINAL_LIST = ",".join(f'$I2="{s}"' for s in sorted(st.TERMINAL))

THIN = Side(style="thin", color=MID_GRAY)


def _font(**kw: Any) -> Font:
    return Font(name="Calibri", **({"size": 11} | kw))


def _local(dt: datetime | None, tz: ZoneInfo) -> datetime | None:
    return as_utc(dt).astimezone(tz).replace(tzinfo=None) if dt else None


def _write_cell(ws, row: int, col: int, value: Any, kind: str) -> None:  # noqa: ANN001
    base = kind.removeprefix("edit-")
    cell = ws.cell(row=row, column=col)
    if base == "url":
        url = safe_url(value) if value else None
        cell.value = excel_safe(url) if url else None
        if url:
            cell.hyperlink = url
            cell.font = _font(color="1F4E79", underline="single")
        return
    cell.value = excel_safe(value) if isinstance(value, str) else value
    if base == "date" and value is not None:
        cell.number_format = DATE_FMT
    elif base == "datetime" and value is not None:
        cell.number_format = DATETIME_FMT
    if base == "num" and value is not None:
        cell.number_format = "#,##0"
    if base in {"text", "note", "status"}:
        cell.alignment = Alignment(vertical="top", wrap_text=base == "note")


def _sheet_table(ws, name: str, columns: list[tuple[str, int, str]], rows: list[list[Any]], freeze: str) -> int:  # noqa: ANN001
    for i, (header, width, kind) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=header)
        editable = kind.startswith("edit-")
        cell.fill = PatternFill("solid", fgColor=NAVY_EDIT if editable else NAVY)
        cell.font = _font(bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if editable:
            cell.comment = Comment("Editable: changes are imported by /sync (or `summer-hunter crm import`).", "Summer Hunter")
        ws.column_dimensions[get_column_letter(i)].width = width
    for r, values in enumerate(rows, start=2):
        for c, ((_, _, kind), value) in enumerate(zip(columns, values, strict=True), start=1):
            _write_cell(ws, r, c, value, kind)
    last = max(len(rows) + 1, 2)          # an Excel table needs at least one body row
    ref = f"A1:{get_column_letter(len(columns))}{last}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 32
    ws.sheet_view.zoomScale = 110
    return last


def _version_label(version: int | None, generated: datetime | None, tz: ZoneInfo) -> str | None:
    if not version:
        return None
    when = _local(generated, tz)
    return f"v{version}" + (f" · {when:%Y-%m-%d}" if when else "")


def _comp(repo: Repository, opp_id: int) -> dict[str, Any]:
    fit = repo.latest_fit(opp_id)
    return (fit.compensation or {}) if fit else {}


def _benefit(comp: dict[str, Any], *names: str) -> str | None:
    found = [f"{n.replace('_', ' ')}: {comp['benefits'][n]}" for n in names if n in (comp.get("benefits") or {})]
    return "; ".join(found) or None


def compensation_cells(repo: Repository, o: Opportunity, tz: ZoneInfo) -> list[Any]:
    comp = _comp(repo, o.id)
    sources = sorted({s.source.split(":")[0] for s in o.sources})
    return [o.compensation_status, o.salary_min, o.salary_max, o.salary_currency, o.salary_period, o.salary_basis,
            o.monthly_min_eur, o.living_cost_eur, o.monthly_balance_eur,
            _benefit(comp, "housing_provided", "housing_allowance"), _benefit(comp, "relocation", "flight"),
            _benefit(comp, "visa_support"), o.compensation_confidence, o.compensation_score,
            o.global_priority, o.next_action, o.strategy_category, ", ".join(sources) or None,
            _local(o.last_verified_at, tz).date() if o.last_verified_at else None]


def _recruit(repo: Repository, opp_id: int) -> dict[str, Any]:
    fit = repo.latest_fit(opp_id)
    return (fit.recruitability or {}) if fit else {}


def application_rows(crm: CRM, tz: ZoneInfo) -> list[list[Any]]:
    rows = []
    repo = Repository(crm.s)
    for a in crm.applications():
        o = a.opportunity
        rec = _recruit(repo, o.id)
        rows.append([
            app_code(a.id), o.company_name, o.title, (o.job_family or "").replace("_", " ") or None,
            country_name(o.country_code) if o.country_code else None, o.city, o.fit_score, o.eligibility_status,
            a.status, _local(o.first_discovered_at, tz).date() if o.first_discovered_at else None, a.deadline,
            a.date_applied, _local(a.updated_at, tz), a.next_action, a.next_action_date, _local(a.interview_at, tz),
            a.recruiter_name, a.recruiter_linkedin, o.official_url, o.apply_url or o.official_url,
            _version_label(a.cv_version, a.cv_generated_at, tz), _version_label(a.letter_version, a.letter_generated_at, tz),
            a.compensation, a.internship_start, a.internship_end, a.notes,
            o.career_category, o.recruitability_score, round(o.industrial_priority) if o.industrial_priority is not None else None,
            "; ".join(rec.get("relevant_experience") or []) or None, ", ".join(rec.get("missing_requirements") or []) or None,
            *compensation_cells(repo, o, tz),
            a.version,
        ])
    return rows


def contact_rows(crm: CRM) -> list[list[Any]]:
    return [[contact_code(c.id), c.company, c.name, c.position, c.linkedin_url,
             app_code(c.application_id) if c.application_id else None, c.contact_type, c.first_contact_date,
             c.last_interaction_date, c.interaction_status, c.next_followup_date, c.next_action, c.notes, c.version]
            for c in crm.contacts()]


def opportunity_rows(crm: CRM, tz: ZoneInfo) -> list[list[Any]]:
    from summer_hunter.crm.deadlines import explicit_deadline

    repo = Repository(crm.s)
    tracked = {a.opportunity_id: a for a in crm.applications()}
    relevant = {o.id: o for o in repo.top(500, include_dismissed=True) if (o.family_tier or 0) > 0 or o.notify}
    for oid in tracked:
        relevant.setdefault(oid, crm.s.get(Opportunity, oid))
    rows = []
    for o in sorted(relevant.values(), key=lambda x: -(x.priority_score or 0)):
        linkedin = crm.s.execute(select(SourceRecord.url).where(SourceRecord.opportunity_id == o.id,
                                                               SourceRecord.source.like("linkedin%"))).scalars().first()
        fit = repo.latest_fit(o.id)
        reasons = "; ".join((fit.report.get("reasons") or [])[:3]) if fit else None
        app = tracked.get(o.id)
        deadline = app.deadline if app else explicit_deadline(o.description_text)
        status = f"{o.status} · {o.user_status}" + (f" · {app_code(app.id)}" if app else "")
        rows.append([o.id, o.company_name, o.title, country_name(o.country_code) if o.country_code else None, o.fit_score,
                     o.networking_score, o.eligibility_status,
                     _local(o.first_discovered_at, tz).date() if o.first_discovered_at else None, deadline,
                     o.official_url, linkedin, o.linkedin_post_url, status, reasons, o.career_category,
                     o.recruitability_score, round(o.industrial_priority) if o.industrial_priority is not None else None])
    return rows


def event_rows(crm: CRM, tz: ZoneInfo) -> list[list[Any]]:
    return [[_local(e.created_at, tz), app_code(e.application_id) if e.application_id else None, e.company, e.event_type,
             e.previous_value, e.new_value, e.source, e.notes] for e in crm.events()]


def _status_formatting(ws, last: int) -> None:  # noqa: ANN001
    rng = f"I2:I{last}"
    for status, color in STATUS_FILL.items():
        font = _font(color=WHITE if color == "A6A6A6" else TEXT, bold=status in st.OFFERS)
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$I2="{status}"'], fill=PatternFill("solid", bgColor=color),
                                                       font=font, stopIfTrue=True))
    ws.conditional_formatting.add(f"N2:O{last}", FormulaRule(
        formula=[f'AND($O2<>"",$O2<TODAY(),NOT(OR({TERMINAL_LIST})))'],
        fill=PatternFill("solid", bgColor="FFC7CE"), font=_font(color="9C0006", bold=True)))
    ws.conditional_formatting.add(f"K2:K{last}", FormulaRule(
        formula=['AND($K2<>"",$K2>=TODAY(),$K2<=TODAY()+7)'],
        fill=PatternFill("solid", bgColor="FFEB9C"), font=_font(color="9C5700", bold=True)))
    ws.conditional_formatting.add(f"G2:G{last}", CellIsRule(operator="greaterThan", formula=["90"],
                                                             fill=PatternFill("solid", bgColor="C6EFCE"),
                                                             font=_font(color="006100", bold=True)))


def _validation(ws, col: str, values: list[str], last: int, prompt: str) -> None:  # noqa: ANN001
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Invalid value", error=f"Choose one of: {', '.join(values)}", promptTitle=prompt,
                        prompt="Pick from the list")
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{max(last, 200)}")


def _date_validation(ws, cols: list[str], last: int) -> None:  # noqa: ANN001
    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                        showErrorMessage=True, errorTitle="Invalid date", error="Enter a date (YYYY-MM-DD).")
    ws.add_data_validation(dv)
    for col in cols:
        dv.add(f"{col}2:{col}{max(last, 200)}")


def _kpi(ws, row: int, col: int, label: str, value: Any, number_format: str | None = None) -> None:  # noqa: ANN001
    lab = ws.cell(row=row, column=col, value=label)
    lab.font = _font(color="595959", size=10)
    lab.fill = PatternFill("solid", fgColor=GRAY)
    lab.border = Border(top=THIN, left=THIN)
    val = ws.cell(row=row, column=col + 1, value="n/a" if value is None else value)
    val.font = _font(bold=True, size=13, color=NAVY)
    val.fill = PatternFill("solid", fgColor=GRAY)
    val.alignment = Alignment(horizontal="right")
    val.border = Border(top=THIN, right=THIN)
    if number_format and value is not None:
        val.number_format = number_format


def _data_block(ws, top: int, col: int, title: str, header: tuple[str, str], rows: list[tuple[Any, int]]) -> tuple[int, int, int]:  # noqa: ANN001
    ws.cell(row=top, column=col, value=title).font = _font(bold=True, color=NAVY)
    ws.cell(row=top + 1, column=col, value=header[0]).font = _font(bold=True, color="595959")
    ws.cell(row=top + 1, column=col + 1, value=header[1]).font = _font(bold=True, color="595959")
    data = rows or [("(none yet)", 0)]
    for i, (k, v) in enumerate(data, start=top + 2):
        ws.cell(row=i, column=col, value=excel_safe(str(k)))
        ws.cell(row=i, column=col + 1, value=v)
    return top + 1, top + 1 + len(data), top + len(data) + 4


def _chart(kind: str, title: str, ws, col: int, first: int, last: int, anchor: str, horizontal: bool = False) -> None:  # noqa: ANN001
    chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[kind]()
    chart.title = title
    data = Reference(ws, min_col=col + 1, min_row=first, max_row=last)
    cats = Reference(ws, min_col=col, min_row=first + 1, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 7.2, 12.5
    if kind == "bar":
        chart.type = "bar" if horizontal else "col"
        chart.legend = None
        chart.series[0].graphicalProperties.solidFill = NAVY
        chart.y_axis.majorGridlines = None
        chart.y_axis.delete = False
        chart.x_axis.delete = False
    if kind == "line":
        chart.legend = None
        chart.series[0].graphicalProperties.line.solidFill = NAVY
        chart.y_axis.delete = False
        chart.x_axis.delete = False
    ws.add_chart(chart, anchor)


def _dashboard(ws, m: Metrics, synced: datetime, tz_name: str) -> None:  # noqa: ANN001
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCDEFGHIJKLMNOPQRSTUVWX", [2, 34, 12, 3, 34, 12, 3, 3, 3, 3, 3, 3, 3, 3, 3, 3, 3, 3, 3, 3, 22, 10, 3, 3], strict=True):
        ws.column_dimensions[col].width = width
    ws.merge_cells("B1:P1")
    title = ws["B1"]
    title.value = "SUMMER HUNTER 2027 · APPLICATION CRM"
    title.font = _font(bold=True, size=18, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = Alignment(vertical="center", indent=1)
    for c in range(2, 17):
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 36
    ws["B2"] = f"Last synchronization: {synced:%Y-%m-%d %H:%M} ({tz_name}). Source of truth: SQLite; this file is a mirror."
    ws["B2"].font = _font(italic=True, color="595959", size=9)

    ws["B4"] = "PIPELINE"
    ws["E4"] = "CONVERSION"
    for ref in ("B4", "E4"):
        ws[ref].font = _font(bold=True, color=NAVY, size=11)
    pipeline = [("Opportunities discovered", m.discovered), ("Opportunities shortlisted", m.shortlisted),
                ("Applications submitted", m.submitted), ("Awaiting response", m.awaiting),
                ("Interviews scheduled", m.interviews_scheduled), ("Offers received", m.offers),
                ("Rejections", m.rejections), ("Withdrawn", m.withdrawn), ("Follow-ups required", m.followups)]
    for i, (label, value) in enumerate(pipeline, start=5):
        _kpi(ws, i, 2, label, value)
    conversion = [("Application → response", m.response_rate), ("Application → interview", m.interview_rate),
                  ("Interview → offer", m.interview_to_offer), ("Overall success (offer / application)", m.success_rate)]
    for i, (label, value) in enumerate(conversion, start=5):
        _kpi(ws, i, 5, label, value, "0.0%")
    ws["E10"] = "n/a = nothing to divide by yet. Rates count only submitted applications."
    ws["E10"].font = _font(italic=True, color="7F7F7F", size=9)
    ws["E12"] = "Legend: dark navy headers are read-only; lighter blue headers are editable and imported by /sync."
    ws["E12"].font = _font(italic=True, color="7F7F7F", size=9)

    col = 21   # chart source tables live in column U onwards, out of the reading area
    top = 4
    blocks = {}
    for key, title_txt, header, rows in (
        ("status", "Applications by status", ("Status", "Count"), [(s, n) for s, n in m.by_status]),
        ("country", "Applications by country", ("Country", "Count"), m.by_country),
        ("family", "Applications by job family", ("Job family", "Count"), m.by_family),
        ("week", "Applications submitted per week", ("Week of", "Submitted"), m.by_week),
        ("funnel", "Application funnel", ("Stage", "Applications"), m.funnel),
    ):
        first, last, top = _data_block(ws, top, col, title_txt, header, rows)
        blocks[key] = (first, last)
    _chart("bar", "Applications by status", ws, col, *blocks["status"], "B16", horizontal=True)
    _chart("bar", "Applications by country", ws, col, *blocks["country"], "H16")
    _chart("bar", "Applications by target job family", ws, col, *blocks["family"], "B32", horizontal=True)
    _chart("line", "Applications submitted over time", ws, col, *blocks["week"], "H32")
    _chart("bar", "Application funnel", ws, col, *blocks["funnel"], "B48", horizontal=True)


APP_EDITABLE = ["status", "date_applied", "next_action", "next_action_date", "interview_at", "notes"]
CONTACT_EDITABLE = ["interaction_status", "next_followup_date", "next_action", "notes"]


def snapshot(crm: CRM) -> dict[str, dict[str, Any]]:
    """Editable values as exported: the common ancestor used to merge your edits on import."""
    from summer_hunter.crm.excel_import import display

    snap = {app_code(a.id): {k: display(getattr(a, k)) for k in APP_EDITABLE} for a in crm.applications()}
    snap.update({contact_code(c.id): {k: display(getattr(c, k)) for k in CONTACT_EDITABLE} for c in crm.contacts()})
    return snap


def _block(ws, top: int, col: int, title: str, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None) -> int:  # noqa: ANN001
    ws.cell(row=top, column=col, value=title).font = _font(bold=True, color=NAVY, size=12)
    for i, h in enumerate(headers):
        c = ws.cell(row=top + 1, column=col + i, value=h)
        c.font = _font(bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
    data = rows or [["(none yet)"] + [None] * (len(headers) - 1)]
    for r, row in enumerate(data, start=top + 2):
        for i, v in enumerate(row):
            cell = ws.cell(row=r, column=col + i, value=excel_safe(v) if isinstance(v, str) else v)
            if isinstance(v, float):
                cell.number_format = "#,##0"
    for i, w in enumerate(widths or []):
        ws.column_dimensions[get_column_letter(col + i)].width = w
    return top + len(data) + 3


def _relevant(crm: CRM) -> list[Opportunity]:
    rows = crm.s.execute(select(Opportunity).where(Opportunity.status == "active")).scalars()
    return [o for o in rows if (o.family_tier or 0) > 0 and o.eligibility_status != "REJECTED"]


def _compensation_sheet(ws, crm: CRM, tz: ZoneInfo) -> None:  # noqa: ANN001
    """Pay is shown as advertised (range minimum in EUR); unknown pay is listed, never counted as zero."""
    ws.sheet_view.showGridLines = False
    opps = _relevant(crm)
    repo = Repository(crm.s)
    ws["A1"] = "COMPENSATION ANALYSIS"
    ws["A1"].font = _font(bold=True, size=16, color=NAVY)
    ws["A2"] = ("Monthly EUR = minimum of the advertised range, converted at the stored ECB / fallback rate. Balances need a "
                "sourced living-cost row in config/compensation.yaml. Unknown pay is not unpaid.")
    ws["A2"].font = _font(italic=True, color="7F7F7F", size=9)
    head = ["Company", "Role", "Country", "Monthly EUR", "Gross / net", "Status", "Confidence", "Global priority"]
    widths = [24, 44, 14, 12, 10, 22, 11, 12]

    def row(o: Opportunity) -> list[Any]:
        return [o.company_name, o.title, country_name(o.country_code) if o.country_code else None, o.monthly_min_eur,
                o.salary_basis, o.compensation_status, o.compensation_confidence, o.global_priority]
    paid = sorted((o for o in opps if o.monthly_min_eur is not None), key=lambda o: -o.monthly_min_eur)
    top = _block(ws, 4, 1, "Highest-paying internships", head, [row(o) for o in paid[:15]], widths)
    bal = sorted((o for o in opps if o.monthly_balance_eur is not None), key=lambda o: -o.monthly_balance_eur)
    top = _block(ws, top, 1, "Best estimated disposable monthly balance", head[:3] + ["Balance EUR", "Living cost EUR", "Monthly EUR", "Confidence"],
                 [[o.company_name, o.title, country_name(o.country_code), o.monthly_balance_eur, o.living_cost_eur,
                   o.monthly_min_eur, o.compensation_confidence] for o in bal[:15]])
    groups: dict[str, list[Opportunity]] = {}
    for o in opps:
        groups.setdefault(country_name(o.country_code) if o.country_code else "Unknown", []).append(o)

    def summary(items: list[Opportunity]) -> list[Any]:
        amounts = [o.monthly_min_eur for o in items if o.monthly_min_eur is not None]
        return [len(items), sum(o.compensation_status == "PAID_CONFIRMED" for o in items),
                sum(o.compensation_status in (None, "COMPENSATION_UNKNOWN") for o in items),
                round(sorted(amounts)[len(amounts) // 2], 0) if amounts else None]
    country_rows = sorted(([k, *summary(v)] for k, v in groups.items()), key=lambda r: -r[1])
    top = _block(ws, top, 1, "Compensation by country", ["Country", "Relevant", "Paid confirmed", "Unknown", "Median monthly EUR"],
                 country_rows[:20])
    fam: dict[str, list[Opportunity]] = {}
    for o in opps:
        fam.setdefault(o.career_category or "Other", []).append(o)
    top = _block(ws, top, 1, "Compensation by career family", ["Category", "Relevant", "Paid confirmed", "Unknown", "Median monthly EUR"],
                 sorted(([k, *summary(v)] for k, v in fam.items()), key=lambda r: -r[1]))
    statuses = ["PAID_CONFIRMED", "PAID_AMOUNT_UNKNOWN", "MIXED_OR_CONDITIONAL", "UNPAID_CONFIRMED", "COMPENSATION_UNKNOWN"]
    status_top = top
    top = _block(ws, top, 1, "Paid versus unpaid", ["Status", "Opportunities"],
                 [[s, sum((o.compensation_status or "COMPENSATION_UNKNOWN") == s for o in opps)] for s in statuses])
    chart = BarChart()
    chart.type, chart.title, chart.legend = "bar", "Paid versus unpaid (relevant opportunities)", None
    chart.add_data(Reference(ws, min_col=2, min_row=status_top + 1, max_row=status_top + 1 + len(statuses)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=status_top + 2, max_row=status_top + 1 + len(statuses)))
    chart.series[0].graphicalProperties.solidFill = NAVY
    chart.height, chart.width = 7, 14
    ws.add_chart(chart, "I4")
    unknown = sorted((o for o in opps if o.compensation_status in (None, "COMPENSATION_UNKNOWN", "PAID_AMOUNT_UNKNOWN")),
                     key=lambda o: -(o.global_priority or 0))
    top = _block(ws, top, 1, "Relevant opportunities with unknown pay (verify before applying)", head,
                 [row(o) for o in unknown[:25]])
    with_benefit = []
    for o in opps:
        comp = _comp(repo, o.id)
        ben = comp.get("benefits") or {}
        if ben:
            with_benefit.append([o.company_name, o.title, country_name(o.country_code),
                                 "yes" if {"housing_provided", "housing_allowance"} & set(ben) else None,
                                 "yes" if {"relocation", "flight"} & set(ben) else None, "yes" if "visa_support" in ben else None])
    _block(ws, top, 1, "Housing / relocation / visa support stated", ["Company", "Role", "Country", "Housing", "Relocation", "Visa"],
           with_benefit[:30])


def _performance_sheet(ws, crm: CRM) -> None:  # noqa: ANN001
    from summer_hunter.engine import analytics
    from summer_hunter.engine.regions import REGIONS

    p = analytics.compute(crm.s, REGIONS)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "HUNTER PERFORMANCE"
    ws["A1"].font = _font(bold=True, size=16, color=NAVY)
    ws["A2"] = f"Rates show their sample; under {analytics.SMALL_SAMPLE} they are indicative only. No hiring probability is inferred."
    ws["A2"].font = _font(italic=True, color="7F7F7F", size=9)
    rows = [["Opportunities discovered", p.discovered], ["Relevant (target family, not rejected)", p.relevant],
            ["Confirmed paid", p.paid_confirmed], ["Unknown compensation", p.compensation_unknown],
            ["Ineligible (rejected)", p.ineligible], ["Duplicate listings merged", p.duplicates_merged],
            ["Recruiter / manager contacts", p.contacts], ["Applications created", p.applications_created],
            ["Applications submitted", p.applications_submitted], ["Reached interview", p.interviews], ["Offers", p.offers],
            ["Precision@10 (your ratings)", p.precision_at_10.label()], ["Actionability rate", p.actionability.label()],
            ["Paid opportunity rate", p.paid_rate.label()], ["Compensation coverage", p.compensation_coverage.label()],
            ["LLM calls / cost (USD)", f"{p.llm_calls} / {p.llm_cost_usd:.2f}"]]
    rows += [[f"Relevant in {k.replace('_', ' ')}", v] for k, v in p.by_region.items()]
    top = _block(ws, 4, 1, "Discovery and conversion", ["Metric", "Value"], rows, [42, 26])
    _block(ws, top, 1, "Source reliability", ["Source", "Last status", "Runs", "Consecutive failures", "Last items"],
           [[s["source"], s["status"], s["runs"], s["failures"], s["last_items"]] for s in p.sources], [42, 14, 8, 12, 10])


def build_workbook(crm: CRM, tz_name: str, export_id: str | None = None) -> Workbook:
    repo = Repository(crm.s)
    ids = [i for (i,) in crm.s.execute(select(Opportunity.id).where(Opportunity.status == "active"))]
    repo.prefetch_fits(ids + [a.opportunity_id for a in crm.applications()])
    try:
        return _build(crm, tz_name, export_id)
    finally:
        repo.clear_fit_cache()


def _build(crm: CRM, tz_name: str, export_id: str | None = None) -> Workbook:
    tz = ZoneInfo(tz_name)
    synced = datetime.now(tz)
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    _dashboard(dash, compute(crm), synced, tz_name)

    apps = wb.create_sheet("Applications")
    last = _sheet_table(apps, "Applications", APPLICATION_COLUMNS, application_rows(crm, tz), "D2")
    _status_formatting(apps, last)
    _validation(apps, "I", st.STATUSES, last, "Application status")
    _date_validation(apps, ["L", "O", "P"], last)
    apps.column_dimensions[get_column_letter(len(APPLICATION_COLUMNS))].hidden = True   # Row Version

    net = wb.create_sheet("Networking")
    last_c = _sheet_table(net, "Networking", CONTACT_COLUMNS, contact_rows(crm), "D2")
    _validation(net, "J", st.INTERACTION_STATUSES, last_c, "Interaction status")
    _date_validation(net, ["K"], last_c)
    net.column_dimensions["N"].hidden = True

    opps = wb.create_sheet("Opportunities")
    _sheet_table(opps, "Opportunities", OPPORTUNITY_COLUMNS, opportunity_rows(crm, tz), "D2")

    events = wb.create_sheet("Activity Log")
    _sheet_table(events, "ActivityLog", EVENT_COLUMNS, event_rows(crm, tz), "A2")

    _compensation_sheet(wb.create_sheet("Compensation Analysis"), crm, tz)
    _performance_sheet(wb.create_sheet("Hunter Performance"), crm)

    meta = wb.create_sheet("_meta")
    for i, (k, v) in enumerate([("format_version", FORMAT_VERSION), ("export_id", export_id or uuid.uuid4().hex),
                                ("exported_at", synced.isoformat(timespec="seconds")), ("timezone", tz_name)], start=1):
        meta.cell(row=i, column=1, value=k)
        meta.cell(row=i, column=2, value=v)
    meta.sheet_state = "hidden"
    return wb


def export_workbook(crm: CRM, path: Path, tz_name: str, backup_dir: Path | None = None) -> Path:
    """Write the workbook atomically; the previous file is kept in `backup_dir` first."""
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and backup_dir is not None:
        backup_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, backup_dir / f"{path.stem}-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
        for old in sorted(backup_dir.glob(f"{path.stem}-*.xlsx"))[:-20]:
            old.unlink()
    export_id = uuid.uuid4().hex
    repo = Repository(crm.s)
    repo.set_kv(f"crm_export:{export_id}", json.dumps(snapshot(crm)))
    history = [k for k in (repo.get_kv("crm_exports") or "").split(",") if k] + [export_id]
    for old_id in history[:-KEEP_SNAPSHOTS]:
        row = crm.s.get(KeyValue, f"crm_export:{old_id}")
        if row is not None:
            crm.s.delete(row)
    repo.set_kv("crm_exports", ",".join(history[-KEEP_SNAPSHOTS:]))
    tmp = path.with_suffix(".tmp.xlsx")
    build_workbook(crm, tz_name, export_id).save(tmp)
    tmp.replace(path)
    return path

