"""Application CRM: records, statuses, reminders, metrics, Excel round trip and Telegram integration."""

from __future__ import annotations

import json
from datetime import date, datetime, timedelta, timezone

import httpx
import pytest
from openpyxl import load_workbook

from summer_hunter.bot.handlers import Bot
from summer_hunter.crm import followups, metrics, sync
from summer_hunter.crm import statuses as st
from summer_hunter.crm.deadlines import explicit_deadline
from summer_hunter.crm.excel_export import FILENAME, build_workbook, export_workbook
from summer_hunter.crm.excel_import import WorkbookError, import_workbook
from summer_hunter.crm.service import CRM
from summer_hunter.crm.validation import ValidationError, excel_safe
from summer_hunter.database.models import Application, CRMEvent, Opportunity
from summer_hunter.database.repository import Repository
from summer_hunter.integrations.anthropic_client import NullProvider
from summer_hunter.integrations.telegram_client import TelegramClient

TOKEN = "123456789:AAbbCCddEEffGGhhIIjjKKllMMnnOOppQQ"
CHAT = "424242"


def make_opp(session, n: int = 1, **fields) -> Opportunity:  # noqa: ANN001
    data = dict(canonical_key=f"k{n}", company_name=f"Company {n}", normalized_company=f"company {n}",
                title=f"Solutions Engineer Intern {n}", normalized_title="solutions engineer intern",
                official_url=f"https://jobs.example.com/{n}", apply_url=f"https://jobs.example.com/{n}/apply",
                country_code="DE", city="Berlin", job_family="technical_sales", family_tier=1, fit_score=80,
                networking_score=40, priority_score=70.0, eligibility_status="ELIGIBLE", status="active",
                description_text="A great internship.")
    data.update(fields)
    opp = Opportunity(**data)
    session.add(opp)
    session.flush()
    return opp


@pytest.fixture
def crm(sessions, settings):
    with sessions() as session:
        yield CRM(session, settings.data_dir, settings.timezone)


# 1 ---------------------------------------------------------------------------------------------
def test_track_creates_application_linked_to_opportunity(crm):
    opp = make_opp(crm.s, description_text="Application deadline: 15 November 2026. Apply now.")
    app, created = crm.track(opp, "telegram")
    assert created and app.status == "TO_APPLY" and app.opportunity_id == opp.id
    assert app.deadline == date(2026, 11, 15) and app.deadline_source == "posting"
    assert opp.user_status == "saved"
    assert [e.event_type for e in crm.events(app.id)] == ["application_created"]


# 2 ---------------------------------------------------------------------------------------------
def test_duplicate_tracking_returns_the_same_record(crm):
    opp = make_opp(crm.s)
    first, _ = crm.track(opp, "telegram")
    second, created = crm.track(opp, "miniapp")
    assert not created and first.id == second.id and len(crm.applications()) == 1


# 3 ---------------------------------------------------------------------------------------------
def test_status_update_logs_history_and_tracks_depth(crm):
    app, _ = crm.track(make_opp(crm.s), "cli")
    crm.set_status(app, "PREPARING", "telegram")
    assert app.status == "PREPARING" and app.version == 2
    ev = crm.events(app.id)[-1]
    assert (ev.event_type, ev.previous_value, ev.new_value, ev.source) == ("status_changed", "TO_APPLY", "PREPARING", "telegram")


# 4 ---------------------------------------------------------------------------------------------
def test_confirmed_submission_sets_date_status_and_hides_from_lists(crm):
    opp = make_opp(crm.s)
    app = crm.confirm_applied(opp, "telegram")
    assert app.status == "APPLIED" and app.date_applied == crm.today()
    assert opp.user_status == "applied"
    assert opp.id not in [o.id for o in Repository(crm.s).top(50)]
    assert "application_submitted" in [e.event_type for e in crm.events(app.id)]


def test_post_application_status_needs_a_confirmed_date(crm):
    app, _ = crm.track(make_opp(crm.s), "cli")
    with pytest.raises(st.TransitionError):
        crm.set_status(app, "APPLIED", "excel")       # never marked applied without confirmation
    with pytest.raises(st.TransitionError):
        crm.set_status(app, "INTERVIEW", "telegram")
    assert app.status == "TO_APPLY"


# 5 ---------------------------------------------------------------------------------------------
def test_recording_an_interview(crm):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    when = datetime(2026, 11, 5, 13, 0, tzinfo=timezone.utc)
    crm.record_interview(app, when, "telegram")
    assert app.status == "INTERVIEW" and app.furthest_status == "INTERVIEW" and app.interview_at == when
    assert "interview_scheduled" in [e.event_type for e in crm.events(app.id)]


# 6 ---------------------------------------------------------------------------------------------
def test_generated_workbook_structure_and_charts(crm, tmp_path):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    crm.add_contact(app.opportunity.company_name, "Jane Recruiter", "telegram", position="Talent", application=app,
                    linkedin_url="https://www.linkedin.com/in/jane", contact_type="Recruiter")
    path = export_workbook(crm, tmp_path / FILENAME, "Europe/Paris")
    wb = load_workbook(path)
    assert wb.sheetnames[:5] == ["Dashboard", "Applications", "Networking", "Opportunities", "Activity Log"]
    assert wb["_meta"].sheet_state == "hidden"
    apps = wb["Applications"]
    headers = [c.value for c in apps[1]]
    assert headers[0] == "Application ID" and headers[25] == "Notes" and headers[-1] == "Row Version" and "Recruitability Score" in headers
    assert apps["A2"].value == "APP-0001" and apps["I2"].value == "APPLIED"
    assert apps["S2"].hyperlink is not None and apps.freeze_panes == "D2"
    assert "Applications" in apps.tables
    assert any("I2" in str(dv.sqref) for dv in apps.data_validations.dataValidation)
    assert len(wb["Dashboard"]._charts) == 5
    assert wb["Networking"]["C2"].value == "Jane Recruiter"


# 7 ---------------------------------------------------------------------------------------------
def test_import_applies_only_authorized_edits(crm, tmp_path):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    path = export_workbook(crm, tmp_path / FILENAME, "Europe/Paris")
    wb = load_workbook(path)
    ws = wb["Applications"]
    ws["I2"] = "INTERVIEW"
    ws["N2"] = "Prepare case study"
    ws["O2"] = datetime(2026, 11, 2)
    ws["Z2"] = "Called by recruiter"
    ws["G2"] = 12                          # read-only score: ignored
    ws["S2"] = "https://evil.example.com"  # read-only URL: ignored
    wb.save(path)
    report = import_workbook(crm, path)
    assert not report.errors and not report.conflicts
    assert app.status == "INTERVIEW" and app.next_action == "Prepare case study"
    assert app.next_action_date == date(2026, 11, 2) and app.notes == "Called by recruiter"
    assert app.opportunity.fit_score == 80 and app.opportunity.official_url.startswith("https://jobs.example.com")


def test_import_is_idempotent(crm, tmp_path):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    path = export_workbook(crm, tmp_path / FILENAME, "Europe/Paris")
    wb = load_workbook(path)
    wb["Applications"]["Z2"] = "first note"
    wb.save(path)
    import_workbook(crm, path)
    before = len(crm.events())
    again = import_workbook(crm, path)
    assert not again.applied and not again.conflicts and len(crm.events()) == before
    assert app.notes == "first note"


def test_import_takes_a_backup_before_changing_anything(crm, tmp_path):
    crm.confirm_applied(make_opp(crm.s), "cli")
    path = export_workbook(crm, tmp_path / FILENAME, "Europe/Paris")
    wb = load_workbook(path)
    wb["Applications"]["Z2"] = "note"
    wb.save(path)
    calls = []
    import_workbook(crm, path, backup=lambda: calls.append(1) or tmp_path / "b.db")
    assert calls == [1]


# 8 ---------------------------------------------------------------------------------------------
def test_invalid_status_and_dates_are_rejected(crm, tmp_path):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    with pytest.raises(st.TransitionError):
        crm.set_status(app, "HIRED", "telegram")
    path = export_workbook(crm, tmp_path / FILENAME, "Europe/Paris")
    wb = load_workbook(path)
    wb["Applications"]["I2"] = "HIRED"
    wb["Applications"]["O2"] = "next tuesday"
    wb.save(path)
    report = import_workbook(crm, path)
    assert report.errors and not report.applied and app.status == "APPLIED"


def test_impossible_transitions_are_refused(crm):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    with pytest.raises(st.TransitionError):
        crm.set_status(app, "TO_APPLY", "telegram")        # cannot un-submit
    crm.set_status(app, "REJECTED", "telegram")
    with pytest.raises(st.TransitionError):
        crm.set_status(app, "INTERVIEW", "telegram")       # final status
    with pytest.raises(ValidationError):
        crm.update_fields(app, {"date_applied": crm.today() + timedelta(days=3)}, "cli")


# 9 ---------------------------------------------------------------------------------------------
def test_conflicting_edit_is_reported_not_applied(crm, tmp_path):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    path = export_workbook(crm, tmp_path / FILENAME, "Europe/Paris")
    crm.update_fields(app, {"notes": "changed in Telegram"}, "telegram")   # database moves on
    wb = load_workbook(path)
    wb["Applications"]["Z2"] = "changed in Excel"
    wb.save(path)
    report = import_workbook(crm, path)
    assert report.conflicts and report.conflicts[0].field == "notes"
    assert app.notes == "changed in Telegram"
    forced = import_workbook(crm, path, prefer_excel=True)
    assert forced.applied and app.notes == "changed in Excel"


def test_stale_workbook_without_edits_is_not_a_conflict(crm, tmp_path):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    path = export_workbook(crm, tmp_path / FILENAME, "Europe/Paris")
    crm.update_fields(app, {"notes": "newer note"}, "telegram")
    report = import_workbook(crm, path)
    assert not report.conflicts and not report.applied and app.notes == "newer note"


# 10 --------------------------------------------------------------------------------------------
def test_follow_up_reminders_are_never_duplicated(crm):
    app = crm.confirm_applied(make_opp(crm.s), "cli", applied_on=date(2026, 10, 1))
    crm.s.flush()
    for ev in crm.events(app.id):
        ev.created_at = datetime(2026, 10, 1, 8, tzinfo=timezone.utc)
    today = date(2026, 10, 20)
    first = followups.pending(crm, today)
    assert [r.kind for r in first] == ["followup_application"]
    followups.mark_sent(crm, first)
    assert followups.pending(crm, today) == []
    assert followups.pending(crm, date(2026, 10, 1)) == []          # not before 10 business days
    crm.set_status(app, "REJECTED", "telegram")
    assert [r for r in followups.all_suggestions(crm, today) if r.kind == "followup_application"] == []


def test_deadline_reminders_7_3_1_without_repeats(crm):
    app, _ = crm.track(make_opp(crm.s), "cli")
    crm.update_fields(app, {"deadline": date(2026, 11, 10)}, "cli")
    assert followups.pending(crm, date(2026, 11, 1)) == []                # 9 days left
    seven = followups.pending(crm, date(2026, 11, 4))
    assert [r.key.rsplit(":", 1)[1] for r in seven] == ["7"]
    followups.mark_sent(crm, seven)
    assert followups.pending(crm, date(2026, 11, 5)) == []
    three = followups.pending(crm, date(2026, 11, 8))                     # 2 days left: the 3-day one
    assert [r.key.rsplit(":", 1)[1] for r in three] == ["3"]
    followups.mark_sent(crm, three)
    assert [r.key.rsplit(":", 1)[1] for r in followups.pending(crm, date(2026, 11, 9))] == ["1"]
    assert followups.pending(crm, date(2026, 11, 11)) == []               # past deadline: nothing


def test_unknown_deadline_is_never_invented():
    assert explicit_deadline("Apply soon! Deadline in November.") is None
    assert explicit_deadline("Applications close on 3 Dec 2026") == date(2026, 12, 3)
    assert explicit_deadline("Apply by January 9, 2027.") == date(2027, 1, 9)


# 11 --------------------------------------------------------------------------------------------
def test_history_is_append_only(crm):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    crm.update_fields(app, {"notes": "one"}, "cli")
    crm.update_fields(app, {"notes": "two"}, "cli")
    notes = [(e.previous_value, e.new_value) for e in crm.events(app.id) if e.notes == "notes"]
    assert notes == [(None, "one"), ("one", "two")]
    assert crm.s.query(CRMEvent).count() == len(crm.events())


def test_legacy_applied_rows_are_migrated(sessions, settings):
    from summer_hunter.crm.migrate import migrate_legacy_applications

    with sessions() as session:
        opp = make_opp(session)
        session.add(Application(opportunity_id=opp.id, applied_at=datetime(2026, 10, 20, 9, tzinfo=timezone.utc)))
        session.flush()
        assert migrate_legacy_applications(session) == 1
        assert migrate_legacy_applications(session) == 0
        app = CRM(session).application_for(opp.id)
        assert app.status == "APPLIED" and app.date_applied == date(2026, 10, 20)


# 12 --------------------------------------------------------------------------------------------
def test_empty_crm_exports_and_computes(crm, tmp_path):
    m = metrics.compute(crm)
    assert m.submitted == 0 and m.response_rate is None and m.success_rate is None
    wb = build_workbook(crm, "Europe/Paris")
    assert wb["Applications"]["A2"].value is None and wb["Dashboard"]["C7"].value == 0
    import_workbook(crm, export_workbook(crm, tmp_path / FILENAME, "Europe/Paris"))


# 13 --------------------------------------------------------------------------------------------
def test_dashboard_metrics_from_real_records(crm):
    apps = [crm.confirm_applied(make_opp(crm.s, n), "cli") for n in range(1, 5)]
    crm.track(make_opp(crm.s, 9), "cli")
    crm.set_status(apps[0], "REJECTED", "cli")
    crm.record_interview(apps[1], datetime(2030, 1, 5, 9, tzinfo=timezone.utc), "cli")
    crm.set_status(apps[1], "OFFER_RECEIVED", "cli")
    crm.record_interview(apps[2], datetime(2030, 1, 6, 9, tzinfo=timezone.utc), "cli")
    crm.set_status(apps[2], "REJECTED", "cli")
    m = metrics.compute(crm)
    assert (m.tracked, m.submitted, m.awaiting, m.rejections, m.offers) == (5, 4, 1, 2, 1)
    assert m.responded == 3 and m.reached_interview == 2
    assert m.response_rate == 0.75 and m.interview_rate == 0.5 and m.interview_to_offer == 0.5 and m.success_rate == 0.25
    assert m.interviews_scheduled == 1
    assert dict(m.funnel)["Submitted"] == 4


# 14 --------------------------------------------------------------------------------------------
def test_formula_injection_is_neutralized(crm, tmp_path):
    opp = make_opp(crm.s, company_name='=HYPERLINK("http://evil","x")', title="+cmd|' /C calc'!A0")
    app, _ = crm.track(opp, "cli")
    crm.update_fields(app, {"notes": "=1+1 then @SUM(A1)"}, "cli")
    path = export_workbook(crm, tmp_path / FILENAME, "Europe/Paris")
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if ws.title != "Dashboard" or cell.column < 20:
                    assert cell.data_type != "f", f"{ws.title}!{cell.coordinate} is a formula"
    assert wb["Applications"]["B2"].value.startswith("'=")
    assert excel_safe("-5") == "'-5" and excel_safe("plain") == "plain"
    # Round trip: the neutralized note comes back unchanged, and a typed formula is refused.
    assert not import_workbook(crm, path).applied and app.notes == "=1+1 then @SUM(A1)"
    wb["Applications"]["N2"] = "=SUM(1,2)"
    wb.save(path)
    report = import_workbook(crm, path)
    assert report.errors and app.next_action is None


def test_macro_or_foreign_files_are_refused(crm, tmp_path):
    bad = tmp_path / "x.xlsm"
    bad.write_bytes(b"not a workbook")
    with pytest.raises(WorkbookError):
        import_workbook(crm, bad)
    other = tmp_path / "other.xlsx"
    from openpyxl import Workbook

    Workbook().save(other)
    with pytest.raises(WorkbookError):
        import_workbook(crm, other)


# 15 --------------------------------------------------------------------------------------------
class FakeTelegram:
    def __init__(self) -> None:
        self.sent: list[dict] = []
        self.docs: list[str] = []
        self.answers: list[str] = []

    def __call__(self, request: httpx.Request) -> httpx.Response:
        method = request.url.path.rsplit("/", 1)[-1]
        if method == "sendDocument":
            self.docs.append(method)
            return httpx.Response(200, json={"ok": True, "result": {"message_id": 99}})
        body = json.loads(request.content or b"{}")
        if method == "sendMessage":
            self.sent.append(body)
            return httpx.Response(200, json={"ok": True, "result": {"message_id": len(self.sent) + 1}})
        if method == "answerCallbackQuery":
            self.answers.append(body.get("text", ""))
        return httpx.Response(200, json={"ok": True, "result": True})


@pytest.fixture
def bot(settings, sessions):
    settings.telegram_bot_token, settings.telegram_chat_id = TOKEN, CHAT
    with sessions() as session:
        make_opp(session, 1)
        make_opp(session, 2)
        session.commit()
    fake = FakeTelegram()
    return Bot(settings, sessions, TelegramClient(TOKEN, transport=httpx.MockTransport(fake)), NullProvider()), fake


def _press(bot: Bot, data: str) -> None:
    bot.handle_update({"update_id": 2, "callback_query": {"id": "cb", "data": data,
                                                          "message": {"message_id": 5, "chat": {"id": int(CHAT)}}}})


def _say(bot: Bot, text: str) -> None:
    bot.handle_update({"update_id": 1, "message": {"chat": {"id": int(CHAT)}, "text": text}})


def test_telegram_track_and_confirmed_applied(bot, sessions):
    b, fake = bot
    _press(b, "tr:1")
    _press(b, "ap:2")
    with sessions() as session:
        crm = CRM(session)
        assert crm.application_for(1).status == "TO_APPLY"
        assert crm.application_for(2) is None                   # Applied only asks for confirmation
    assert any("Confirmer" in m["text"] for m in fake.sent)
    _press(b, "apc:2")
    with sessions() as session:
        app = CRM(session).application_for(2)
        assert app.status == "APPLIED" and app.date_applied is not None
        assert [f.action for f in Repository(session).feedback_rows()] == ["applied"]
    _press(b, "cs:2:INTERVIEW")
    _press(b, "cs:2:TO_APPLY")                                  # refused, history intact
    with sessions() as session:
        assert CRM(session).application_for(2).status == "INTERVIEW"


def test_telegram_crm_commands(bot, settings):
    b, fake = bot
    _press(b, "tr:1")
    for cmd in ("/crm", "/applications", "/interviews", "/followups", "/deadlines", "/applied", "/offers",
                "/company company 1", "/deadline APP-1 2030-01-01", "/note APP-1 call Monday", "/setstatus APP-1 PREPARING",
                "/export", "/sync", "/top10", "/stats", "/help"):
        _say(b, cmd)
    texts = "\n".join(m["text"] for m in fake.sent)
    for expected in ("CRM candidatures", "Candidatures actives", "Dates limites", "Top 10", "Stats", "/followups"):
        assert expected in texts
    assert fake.docs and sync.workbook_path(settings).exists()
    assert "Import Excel" in texts
    assert TOKEN not in texts
    cards = [m for m in fake.sent if m.get("reply_markup")]
    assert all(len(btn.get("callback_data", "")) <= 64 for m in cards for row in m["reply_markup"]["inline_keyboard"] for btn in row)


def test_card_buttons_include_crm_actions(sessions):
    from summer_hunter.notifications.formatter import buttons

    with sessions() as session:
        opp = make_opp(session)
        labels = [b["text"] for row in buttons(opp, None) for b in row]
    for label in ("Save", "Track", "Applied", "Dismiss"):
        assert label in labels


# Regressions found in the independent review --------------------------------------------------
def test_duplicate_ids_in_workbook_are_rejected(crm, tmp_path):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    path = export_workbook(crm, tmp_path / FILENAME, "Europe/Paris")
    wb = load_workbook(path)
    ws = wb["Applications"]
    for col in range(1, ws.max_column + 1):
        ws.cell(row=3, column=col, value=ws.cell(row=2, column=col).value)
    ws["Z2"], ws["Z3"] = "version A", "version B"
    wb.save(path)
    report = import_workbook(crm, path)
    assert report.errors and not report.applied and app.notes is None


def test_date_applied_in_excel_is_a_confirmed_submission(crm, tmp_path, settings):
    app, _ = crm.track(make_opp(crm.s), "cli")
    path = export_workbook(crm, tmp_path / FILENAME, "Europe/Paris")
    wb = load_workbook(path)
    wb["Applications"]["L2"] = datetime(2026, 10, 20)
    wb.save(path)
    import_workbook(crm, path)
    assert app.status == "APPLIED" and app.date_applied == date(2026, 10, 20)
    assert app.submitted_folder and (settings.data_dir / "applications").exists()
    assert "application_submitted" in [e.event_type for e in crm.events(app.id)]


def test_backward_pipeline_moves_need_force(crm):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    crm.set_status(app, "FINAL_INTERVIEW", "cli")
    with pytest.raises(st.TransitionError):
        crm.set_status(app, "APPLIED", "cli")
    crm.set_status(app, "INTERVIEW", "cli", force=True)
    assert app.status == "INTERVIEW" and app.furthest_status == "FINAL_INTERVIEW"


def test_import_uses_the_exported_timezone(crm, tmp_path):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    crm.record_interview(app, datetime(2026, 11, 5, 13, 0, tzinfo=timezone.utc), "cli")
    path = export_workbook(crm, tmp_path / FILENAME, "Europe/Paris")
    crm.tz = "UTC"                                   # configuration changed after the export
    wb = load_workbook(path)
    wb["Applications"]["Z2"] = "unrelated note"
    wb.save(path)
    report = import_workbook(crm, path)
    assert report.applied == ["APP-0001 notes: unrelated note"]
    assert app.interview_at.replace(tzinfo=timezone.utc) == datetime(2026, 11, 5, 13, 0, tzinfo=timezone.utc)


def test_save_or_dismiss_never_undo_a_submission(crm):
    opp = make_opp(crm.s)
    crm.confirm_applied(opp, "cli")
    repo = Repository(crm.s)
    repo.record_feedback(opp, "dismiss")
    repo.record_feedback(opp, "save")
    assert opp.user_status == "applied"


def test_backdated_application_follow_up_uses_the_real_date(crm):
    app = crm.confirm_applied(make_opp(crm.s), "migration", applied_on=date(2026, 9, 1))
    assert app.date_applied == date(2026, 9, 1)
    assert [r.kind for r in followups.pending(crm, crm.today())] == ["followup_application"]


def test_submission_without_documents_is_frozen_empty(crm):
    app = crm.confirm_applied(make_opp(crm.s), "cli")
    from pathlib import Path

    manifest = Path(app.submitted_folder) / "MANIFEST.txt"
    assert "none existed" in manifest.read_text()
    cv_dir = crm.data_dir / "cv" / str(app.opportunity_id)
    cv_dir.mkdir(parents=True)
    (cv_dir / "late.pdf").write_bytes(b"%PDF late")
    crm.sync_documents(app)
    assert app.cv_version == 0                      # a CV generated after submission is not the one sent
