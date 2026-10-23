"""Industry & robotics track: taxonomy, recruitability, eligibility, ranking, Telegram, CRM, discovery."""

from __future__ import annotations

import httpx
import pytest
import yaml

from summer_hunter.bot.handlers import Bot
from summer_hunter.crm.excel_export import FILENAME, export_workbook
from summer_hunter.crm.service import CRM
from summer_hunter.engine import recruitability, requirements
from summer_hunter.engine.discovery import plan
from summer_hunter.integrations.anthropic_client import NullProvider
from summer_hunter.integrations.telegram_client import TelegramClient
from summer_hunter.models.opportunity import RawOpportunity
from summer_hunter.notifications import formatter
from summer_hunter.pipeline import CycleStats
from summer_hunter.utils.settings import PROJECT_ROOT
from tests.integration.test_crm import CHAT, TOKEN, FakeTelegram

DATES = "The internship runs from April 6 to August 28, 2027 (5 months). "

MAINTENANCE = (DATES + "Join our MRO team at an airline. You will analyze maintenance data, monitor maintenance KPIs, support "
               "preventive maintenance planning and drive process improvement with engineering and operations stakeholders. "
               "Requirements: engineering student in mechanical or industrial engineering, strong data analysis skills (Excel, "
               "Python), good communication in English. SAP PM is a plus.")
INDUSTRIAL_PM = (DATES + "You will support industrial projects in our manufacturing plant: project coordination, industrial "
                 "planning, budget monitoring, project reporting and cross-functional collaboration with production and quality. "
                 "Profile: engineering school student in industrial engineering, strong organization, Excel. Lean is a plus.")
ROBOTICS_DEPLOY = (DATES + "Help deploy autonomous mobile robots at customer warehouses: customer onboarding, deployment "
                   "coordination, fleet performance monitoring and technical implementation with the engineering team. You have: "
                   "engineering studies in robotics, mechatronics or industrial engineering, project coordination, customer-facing "
                   "communication. Exposure to ROS is a plus.")
RAAS_OPS = (DATES + "Our robots-as-a-service fleet runs in 40 factories. As Robotics Operations intern you will track fleet "
            "performance KPIs, coordinate service delivery and maintenance of robotic systems, and analyse uptime data with "
            "Python and Excel. Industrial or mechanical engineering students welcome.")
ROBOTICS_RESEARCH = (DATES + "Research internship on SLAM and perception. Requirements: strong C++ and ROS2 development "
                     "experience, proven experience with SLAM algorithms and computer vision research.")
PLC_CORE = (DATES + "Automation intern. Requirements: proven PLC programming experience (Siemens TIA Portal) is mandatory; "
            "you will commission production lines.")
VAGUE = DATES + "Great internship in our robotics team. Apply now to join a fast growing company."


@pytest.fixture
def scoring(settings):
    return settings.scoring


def _evaluate(hunter, repo, title, description, location="Munich, Germany", company="ACME Robotics", sid="1"):
    stats = CycleStats()
    raw = RawOpportunity(source="manual", source_id=sid, url=f"https://jobs.example.com/{sid}", title=title,
                         company=company, location_text=location, description_text=description)
    opp = hunter.ingest(repo, raw, stats)
    hunter.evaluate(repo, opp, None, stats)
    return opp


# 1-4: target internships -------------------------------------------------------------------------
@pytest.mark.parametrize(("title", "text", "family"), [
    ("Maintenance Engineering Intern", MAINTENANCE, "maintenance_reliability"),
    ("Industrial Project Management Intern", INDUSTRIAL_PM, "industrial_project_management"),
    ("Robotics Deployment Intern", ROBOTICS_DEPLOY, "robotics_raas"),
    ("Robotics Operations Intern", RAAS_OPS, "robotics_raas"),
])
def test_target_industrial_internships_score_well(hunter, repo, title, text, family):
    opp = _evaluate(hunter, repo, title, text)
    assert opp.job_family == family and opp.career_track == "industry_robotics"
    assert opp.recruitability_score >= 70, opp.recruitability_score
    assert opp.eligibility_status != "REJECTED"
    assert opp.industrial_priority is not None and opp.priority_score == opp.industrial_priority


def test_maintenance_matches_air_france_experience(hunter, repo):
    opp = _evaluate(hunter, repo, "Maintenance Engineering Intern", MAINTENANCE)
    rec = repo.latest_fit(opp.id).recruitability
    assert any("Air France" in r for r in rec["relevant_experience"])
    assert rec["components"]["industry"] == 90 and not rec["blocking_requirements"]


# 5-6: specialized engineering and mandatory mismatches ---------------------------------------------
def test_specialized_robotics_software_is_low_and_fails_eligibility(hunter, repo):
    opp = _evaluate(hunter, repo, "Robotics Software Research Intern", ROBOTICS_RESEARCH)
    assert opp.job_family == "robotics_specialized" and opp.recruitability_score <= 20
    assert opp.eligibility_status == "REJECTED"
    crit = {c["criterion"]: c for c in repo.latest_eligibility(opp.id).criteria}
    assert crit["technical_requirements"]["status"] == "FAIL"


def test_mandatory_skill_mismatch_vs_mere_mention(settings, profile):
    lex = settings.scoring["recruitability"]["skills"]
    core = requirements.extract(PLC_CORE, lex, profile.verified_skills)
    assert [(r.name, r.level, r.core) for r in requirements.blocking(core)] == [("plc programming", "mandatory", True)]
    mention = requirements.extract("You will get exposure to PLCs and CAD drawings on the shop floor.", lex, profile.verified_skills)
    assert {r.level for r in mention} == {"exposure"} and requirements.blocking(mention) == []
    plus = requirements.extract("Requirements: engineering student. Nice to have: ROS2, SolidWorks.", lex, profile.verified_skills)
    assert {r.level for r in plus} == {"preferred"}


def test_section_headings_set_the_requirement_level(settings, profile):
    text = "Basic qualifications:\n- Proficiency in SolidWorks\nPreferred qualifications:\n- Experience with ROS2\n"
    found = {r.name: r.level for r in requirements.extract(text, settings.scoring["recruitability"]["skills"], profile.verified_skills)}
    assert found == {"cad design": "mandatory", "ros / ros2": "preferred"}


# 7: unknown requirements ------------------------------------------------------------------------
def test_unknown_technical_requirements_are_neutral_not_penalized(settings, profile, title_rules):
    fam = title_rules.classify("Robotics Intern")
    rep = recruitability.score("Robotics Intern", VAGUE, fam, profile, settings.scoring)
    assert not rep.technical_requirements_known
    assert rep.components["technical_skills"] == settings.scoring["recruitability"]["unknown_technical_score"]
    assert rep.blocking_requirements == [] and any("not stated" in c for c in rep.concerns)


# 8: recruitability scoring -------------------------------------------------------------------------
def test_recruitability_components_weights_and_honesty(settings, profile, title_rules):
    w = settings.scoring["recruitability"]["weights"]
    assert w == {"direct_experience": 0.30, "technical_skills": 0.25, "academic": 0.20, "transferable": 0.15, "industry": 0.10}
    fam = title_rules.classify("Maintenance Engineering Intern")
    rep = recruitability.score("Maintenance Engineering Intern", MAINTENANCE, fam, profile, settings.scoring)
    expected = sum(rep.components[k] * v for k, v in w.items())
    assert abs(rep.score - expected) <= 1
    assert "not a probability" in rep.note
    bad = recruitability.score("Robotics Software Research Intern", ROBOTICS_RESEARCH,
                               title_rules.classify("Robotics Software Research Intern"), profile, settings.scoring)
    assert bad.score < rep.score - 50 and len(bad.blocking_requirements) >= 2


# 9: existing career fit unchanged for the business track -----------------------------------------
def test_business_track_keeps_its_scoring(hunter, repo):
    from tests.conftest import FILLER

    opp = _evaluate(hunter, repo, "Solutions Engineer Intern", FILLER + DATES, location="Singapore", company="Stripe", sid="9")
    assert opp.job_family == "technical_sales" and opp.career_track == "business_tech"
    fit = repo.latest_fit(opp.id)
    assert opp.priority_score == round(0.75 * opp.fit_score + 0.25 * opp.networking_score, 1)   # unchanged formula
    assert fit.recruitability is not None and opp.recruitability_score is not None             # informative only


def test_site_reliability_stays_software(title_rules):
    assert title_rules.classify("Site Reliability Engineer Intern").id == "software_engineering"
    assert title_rules.classify("Reliability Engineering Intern").id == "maintenance_reliability"
    assert title_rules.classify("Manufacturing Sales Intern").id == "sales_general"
    assert title_rules.classify("Mechanical Design Intern").id == "design_engineering"     # CAD-centred design roles
    assert not title_rules.worth_collecting("Mechanical Design Intern")
    assert title_rules.worth_collecting("Robotics Software Engineer Intern")   # collected, then judged on requirements


# 10-12: hard constraints still apply -------------------------------------------------------------------
def test_geographic_exclusion_applies_to_industrial_roles(hunter, repo):
    opp = _evaluate(hunter, repo, "Maintenance Engineering Intern", MAINTENANCE, location="Toulouse, France", sid="10")
    assert opp.eligibility_status == "REJECTED" and opp.industrial_priority is None and not opp.notify


def test_language_requirement_is_enforced(hunter, repo):
    text = MAINTENANCE + " Fluent German is required for daily work with our technicians."
    opp = _evaluate(hunter, repo, "Maintenance Planning Intern", text, sid="11")
    crit = {c["criterion"]: c["status"] for c in repo.latest_eligibility(opp.id).criteria}
    assert crit["language"] == "FAIL" and opp.eligibility_status == "REJECTED"


def test_dates_outside_the_window_are_rejected(hunter, repo):
    text = MAINTENANCE.replace(DATES, "The internship runs from January 5 to March 27, 2027. ")
    opp = _evaluate(hunter, repo, "Maintenance Engineering Intern", text, sid="12")
    crit = {c["criterion"]: c["status"] for c in repo.latest_eligibility(opp.id).criteria}
    assert crit["dates"] == "FAIL"


# 13: CRM ------------------------------------------------------------------------------------------------
def test_crm_and_excel_show_track_fields(hunter, repo, settings, tmp_path):
    from openpyxl import load_workbook

    opp = _evaluate(hunter, repo, "Maintenance Engineering Intern", MAINTENANCE, sid="13")
    crm = CRM(repo.s, settings.data_dir, settings.timezone)
    crm.track(opp, "telegram")
    wb = load_workbook(export_workbook(crm, tmp_path / FILENAME, "Europe/Paris"))
    ws = wb["Applications"]
    row = {ws.cell(row=1, column=c).value: ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)}
    assert row["Career Category"] == "Maintenance" and row["Recruitability Score"] == opp.recruitability_score
    assert row["Industrial Priority Score"] == round(opp.industrial_priority)
    assert "Air France" in row["Relevant Industrial Experience"]
    assert ws.column_dimensions["AF"].hidden          # Row Version moved to the end
    opps = wb["Opportunities"]
    headers = [c.value for c in opps[1]]
    assert "Recruitability Score" in headers and "Career Category" in headers


# 14: Telegram --------------------------------------------------------------------------------------------
def test_industrial_card_and_commands(hunter, repo, settings, sessions):
    opp = _evaluate(hunter, repo, "Robotics Deployment Intern", ROBOTICS_DEPLOY, sid="14")
    repo.s.commit()
    fit, ev = repo.latest_fit(opp.id), repo.latest_eligibility(opp.id)
    text, buttons = formatter.card(opp, ev.criteria, fit.report, fit.networking, fit.priority, [], None, fit.recruitability)
    assert "INDUSTRIAL INTERNSHIP" in text and "Recruitability:" in text and "not a probability" in text
    assert "Category:</b> INDUSTRY &amp; ROBOTICS · Robotics" in text
    assert {"Save", "Track", "Applied", "Dismiss"} <= {b["text"] for row in buttons for b in row}
    settings.telegram_bot_token, settings.telegram_chat_id = TOKEN, CHAT
    fake = FakeTelegram()
    bot = Bot(settings, sessions, TelegramClient(TOKEN, transport=httpx.MockTransport(fake)), NullProvider())
    for cmd in ("/industry", "/maintenance", "/robotics", "/industrial_projects", "/high_recruitability", "/top10", "/applied"):
        bot.handle_update({"update_id": 1, "message": {"chat": {"id": int(CHAT)}, "text": cmd}})
    texts = "\n".join(m["text"] for m in fake.sent)
    assert "Industry &amp; robotics track" in texts and "Robotics Deployment Intern" in texts
    assert "Recruitability ≥" in texts and "Top 10" in texts


# 15: discovery -----------------------------------------------------------------------------------------
def test_discovery_plan_splits_tracks_and_rotates():
    from collections import Counter

    from summer_hunter.engine.discovery import plan_by_track

    disc = yaml.safe_load((PROJECT_ROOT / "config" / "search_preferences.yaml").read_text())["discovery"]
    first = plan_by_track(disc, 0)
    assert len(first) == disc["max_queries_per_run"]
    assert Counter(t for t, _, _ in first) == {"business_tech": 15, "industry_robotics": 15}      # 50 / 50
    seen = {q for r in range(6) for _, _, q in plan_by_track(disc, r)}
    assert '"Robotics as a Service" internship 2027' in seen
    assert plan(disc, 0) != plan(disc, 1)                                                    # rotation
    assert any(k == "job" and "2027" not in q for _, k, q in first)                           # yearless variants
    lopsided = dict(disc, allocation={"business_tech": 0.2, "industry_robotics": 0.8})
    assert Counter(t for t, _, _ in plan_by_track(lopsided, 0)) == {"business_tech": 6, "industry_robotics": 24}
    one = dict(disc, max_queries_per_run=1)                                                   # ties rotate
    assert {plan_by_track(one, r)[0][0] for r in range(2)} == {"business_tech", "industry_robotics"}


def test_linkedin_collector_uses_the_plan():
    from summer_hunter.collectors.linkedin_public import LinkedInSearchCollector

    c = LinkedInSearchCollector(None, [], [], planned=[("job", '"Robotics Operations" intern 2027')])
    assert c.planned == [("job", '"Robotics Operations" intern 2027')]


def test_alternatives_and_industry_words_are_not_hard_requirements(settings, profile):
    lex = settings.scoring["recruitability"]["skills"]
    alt = requirements.extract("Requirements: programming experience in C++, Java or Python.", lex, profile.verified_skills)
    assert requirements.blocking(alt) == []
    sector = requirements.extract("Required: passion for the semiconductor industry and strong communication.", lex,
                                  profile.verified_skills)
    assert requirements.blocking(sector) == []


def test_review_regressions_on_requirement_levels(settings, profile):
    lex = settings.scoring["recruitability"]["skills"]

    def levels(text):
        return {r.name: r.level for r in requirements.extract(text, lex, profile.verified_skills)}

    assert levels("Minimum qualifications: C++ required, Python is a plus") == {"c++": "mandatory", "python": "preferred"}
    assert [r.name for r in requirements.blocking(requirements.extract("C++ and either Python or SQL required", lex,
                                                                       profile.verified_skills))] == ["c++"]
    duties = "Requirements: Excel required\nResponsibilities: Build advanced computer vision prototypes"
    assert levels(duties)["computer vision"] == "exposure"
    assert levels("Requirements: Strong understanding of customer perception and market research") == {}


def test_review_regressions_on_classification(title_rules):
    assert title_rules.classify("Maintenance Sales Intern").id == "sales_general"
    assert title_rules.classify("Technical Services Sales Intern").id == "sales_general"
    assert title_rules.classify("Product Project Manager Intern").id == "product_management"
    assert title_rules.classify("Customer Perception Research Intern").id != "robotics_specialized"
    assert title_rules.classify("Robotics Engineer Intern").id == "robotics_specialized"


def test_records_scored_before_the_track_are_reevaluated(hunter, repo):
    opp = _evaluate(hunter, repo, "Maintenance Engineering Intern", MAINTENANCE, sid="20")
    assert opp not in repo.opportunities_needing_evaluation()
    opp.career_track = None                       # as left by the previous version
    assert opp in repo.opportunities_needing_evaluation()


def test_comma_list_with_final_or_is_an_alternative(settings, profile):
    lex = settings.scoring["recruitability"]["skills"]
    assert requirements.blocking(requirements.extract("Proficiency in Python, C++, R, or SQL", lex, profile.verified_skills)) == []
    both = requirements.extract("Proficiency in Python and C++, with deep-learning frameworks", lex, profile.verified_skills)
    assert [r.name for r in requirements.blocking(both)] == ["c++"]
