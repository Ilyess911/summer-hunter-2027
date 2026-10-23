"""V3: Africa, compensation intelligence, global priority, strategist, signals, networking, analytics, Telegram."""

from __future__ import annotations

from datetime import date, datetime, timedelta, timezone

import httpx
import pytest
import yaml

from summer_hunter.bot.handlers import Bot
from summer_hunter.crm.service import CRM
from summer_hunter.database.models import FxRate, LinkedInPost, SourceHealth
from summer_hunter.engine import analytics, fx, knowledge, networking_reco, priority, semantic, signals, strategist
from summer_hunter.engine.compensation import PAID_AMOUNT_UNKNOWN, PAID_CONFIRMED, UNPAID_CONFIRMED, assess
from summer_hunter.engine.discovery import plan_by_track
from summer_hunter.engine.fx import FxTable, Rate
from summer_hunter.engine.geo import resolve_location
from summer_hunter.engine.regions import REGIONS, region_of
from summer_hunter.integrations.anthropic_client import NullProvider
from summer_hunter.integrations.telegram_client import TelegramClient
from summer_hunter.models.opportunity import RawOpportunity
from summer_hunter.pipeline import CycleStats, in_backoff
from summer_hunter.utils.settings import PROJECT_ROOT
from tests.integration.test_crm import CHAT, TOKEN, FakeTelegram

D = date(2026, 10, 23)
FX = FxTable({"EUR": Rate("EUR", 1.0, None, "identity"), "USD": Rate("USD", 1.12, D, "ECB euro reference rates"),
              "MAD": Rate("MAD", 11.0, D, "open.er-api.com (ExchangeRate-API)"), "GBP": Rate("GBP", 0.84, D, "ECB"),
              "XOF": Rate("XOF", 655.957, None, "fixed peg")})
DATES = "The internship runs from April 6 to August 28, 2027 (20 weeks). "
CFG = {"hours": {"weeks_per_month": 4.333, "assumed_hours_per_week_for_estimates": 40, "allow_hourly_estimate": True},
       "scoring": {"weights": {"disposable_balance": .30, "benchmark": .25, "housing_relocation": .20, "total_outcome": .15,
                               "confidence": .10},
                   "balance_scale_eur": [-500, 1500], "monthly_scale_eur": [0, 3000], "total_outcome_scale_eur": [-2000, 8000],
                   "confidence_points": {"HIGH": 100, "MEDIUM": 70, "LOW": 40, "UNKNOWN": 0}, "unpaid_score": 0},
       "living_costs": [{"place": "casablanca, ma", "housing": 400, "transport": 40, "food": 200, "other": 60,
                         "source": "test dataset", "date": "2026-10-01"}],
       "benchmarks": []}


def comp(text, country="US", city=None, source="greenhouse:acme", weeks=20, cfg=CFG):  # noqa: ANN001
    return assess(text, country, city, source, weeks, cfg, FX)


# 1-3 Africa and queries ----------------------------------------------------------------------------
def test_africa_is_a_full_region(settings):
    geo = settings.preferences["geography"]
    assert set(geo["priority_countries"]["africa"]) >= {"MA", "TN", "EG", "SN", "CI", "GH", "NG", "KE", "RW", "TZ", "UG", "ZA", "BW", "NA"}
    for place, code in [("Casablanca", "MA"), ("Benguerir", "MA"), ("Lagos", "NG"), ("Kigali", "RW"), ("Gaborone", "BW"),
                        ("Windhoek", "NA"), ("Cape Town", "ZA"), ("Abidjan", "CI")]:
        assert resolve_location(place, None).primary_country == code
        assert region_of(code) == "africa"
    assert not set(REGIONS["africa"]) & set(geo["excluded_countries"])


def test_exclusions_kept_france_israel_and_previous_choices(hunter, repo):
    excluded = set(hunter.settings.preferences["geography"]["excluded_countries"])
    assert {"FR", "IL"} <= excluded and not excluded & {"NG", "KE", "EG", "ZA"}
    stats = CycleStats()
    for sid, loc in (("fr", "Paris, France"), ("il", "Tel Aviv, Israel"), ("ma", "Casablanca, Morocco")):
        opp = hunter.ingest(repo, RawOpportunity(source="manual", source_id=sid, url=f"https://x.test/{sid}", title="Maintenance Engineering Intern",
                                                 company="Acme", location_text=loc, description_text=DATES + "maintenance KPIs"), stats)
        hunter.evaluate(repo, opp, None, stats)
        crit = {c["criterion"]: c["status"] for c in repo.latest_eligibility(opp.id).criteria}
        assert (crit["country"] == "FAIL") == (sid in {"fr", "il"})


def test_french_and_english_african_queries():
    disc = yaml.safe_load((PROJECT_ROOT / "config" / "search_preferences.yaml").read_text())["discovery"]
    phrases = [q for r in range(8) for _, _, q in plan_by_track(disc, r)]
    assert any("stage" in q and ("Maroc" in q or "Casablanca" in q) for q in phrases)
    assert any("internship" in q and ("Nairobi" in q or "Lagos" in q or "South Africa" in q) for q in phrases)


# 4-13 Compensation ---------------------------------------------------------------------------------
def test_paid_detection_and_range_kept():
    r = comp("The hourly pay range for this internship is $29 - $33 per hour, 40 hours per week.")
    assert r.status == PAID_CONFIRMED and r.pay.amount_min == 29 and r.pay.amount_max == 33 and r.pay.currency == "USD"
    assert r.monthly_min < r.monthly_max and r.confidence == "HIGH"


def test_unknown_compensation_is_not_zero_nor_unpaid():
    r = comp("Great team at a multinational, apply now.")
    assert r.status == "COMPENSATION_UNKNOWN" and r.score is None and r.monthly_min_eur is None and r.confidence == "UNKNOWN"
    assert comp("We offer a competitive stipend.").status == PAID_AMOUNT_UNKNOWN
    assert comp("This is an unpaid internship.").status == UNPAID_CONFIRMED
    assert comp("A stipend may be provided depending on your university.").status == "MIXED_OR_CONDITIONAL"


def test_hourly_to_monthly_exact_and_estimated():
    exact = comp("Pay: $25 per hour for 20 hours per week.")
    assert exact.hours_source == "posting" and not exact.monthly_is_estimate
    assert exact.monthly_min == pytest.approx(25 * 20 * 4.333, rel=1e-3)
    est = comp("Pay: $25 per hour.")
    assert est.monthly_is_estimate and est.hours_source == "assumption" and any("assumed" in n for n in est.notes)
    no_est = comp("Pay: $25 per hour.", cfg=CFG | {"hours": {"allow_hourly_estimate": False}})
    assert no_est.monthly_min is None
    assert comp("Monthly stipend of $1,800 for 40 hours per week").pay.period == "month"
    assert comp("Annual salary range: $67,498.00-67,502.00 USD (hourly role)").pay.period == "year"   # implausible hourly


def test_currency_conversion_records_rate_date_source():
    r = comp("Gratification de 6 000 MAD par mois.", country="MA", city="Casablanca")
    assert r.monthly_min_eur == pytest.approx(6000 / 11.0, abs=0.01)
    assert (r.fx_rate, r.fx_date) == (11.0, "2026-10-23") and "er-api" in r.fx_source
    assert comp("Stipend 150 000 FCFA par mois", country="SN").monthly_min_eur == pytest.approx(150000 / 655.957, abs=0.01)
    unknown = comp("Stipend of 900 CHF per month", country="CH")
    assert unknown.monthly_min is not None and unknown.monthly_min_eur is None and any("No exchange rate" in n for n in unknown.notes)


def test_fx_sources_parse_and_pegs_win():
    day, rates = fx.parse_ecb("<Cube time='2026-10-23'><Cube currency='USD' rate='1.1193'/><Cube currency='GBP' rate='0.83728'/></Cube>")
    assert day == date(2026, 10, 23) and rates == {"USD": 1.1193, "GBP": 0.83728}
    day2, rates2 = fx.parse_er({"result": "success", "time_last_update_unix": 1792800000, "rates": {"MAD": 11.0, "XOF": 655.0}})
    assert rates2["MAD"] == 11.0 and day2 is not None
    table = FxTable.from_rows([FxRate(currency="XOF", per_eur=655.0, rate_date=D, source="er")], {"XOF": 655.957})
    assert table.rates["XOF"].per_eur == 655.957 and table.rates["XOF"].source == "fixed peg"


def test_fx_refresh_stores_with_mocked_sources(sessions, settings):
    def handler(request: httpx.Request) -> httpx.Response:
        if "ecb" in str(request.url):
            return httpx.Response(200, text="<Cube time='2026-10-23'><Cube currency='USD' rate='1.12'/></Cube>")
        return httpx.Response(200, json={"result": "success", "time_last_update_unix": 1792800000, "rates": {"USD": 1.13, "MAD": 11}})
    with sessions() as s:
        out = fx.refresh(s, settings.compensation, httpx.Client(transport=httpx.MockTransport(handler)))
        table = fx.load(s, settings.compensation)
    assert out["stored"] == 2 and table.rates["USD"].per_eur == 1.12 and table.rates["MAD"].source.startswith("open.er-api")


def test_gross_and_net_are_recorded_not_mixed():
    assert comp("Salary: £2,000 per month gross", country="GB").pay.basis == "gross"
    assert comp("You will receive €900 net per month", country="DE").pay.basis == "net"
    unknown = comp("Salary: £2,000 per month", country="GB")
    assert unknown.pay.basis == "unknown" and any("Gross or net" in n for n in unknown.notes)


def test_housing_reduces_costs_not_cash():
    with_housing = comp("Gratification : 6 000 MAD par mois. Logement fourni.", country="MA", city="Casablanca")
    without = comp("Gratification : 6 000 MAD par mois.", country="MA", city="Casablanca")
    assert "housing_provided" in with_housing.benefits
    assert with_housing.living_cost_breakdown["housing"] == 0 and with_housing.monthly_min_eur == without.monthly_min_eur
    assert with_housing.monthly_balance_eur - without.monthly_balance_eur == pytest.approx(400)


def test_relocation_benefit_raises_attractiveness():
    base = comp("Pay: $3,000 per month.")
    reloc = comp("Pay: $3,000 per month. Relocation assistance and round-trip flight provided.")
    assert {"relocation", "flight"} <= set(reloc.benefits)
    assert reloc.score > base.score


def test_cost_of_living_only_with_a_sourced_row():
    sourced = comp("Gratification : 6 000 MAD par mois.", country="MA", city="Casablanca")
    assert sourced.living_cost_eur == 700 and "test dataset" in sourced.living_cost_source
    assert sourced.monthly_balance_eur == pytest.approx(6000 / 11 - 700, abs=0.01) and sourced.total_outcome_eur is not None
    missing = comp("Gratification : 6 000 MAD par mois.", country="MA", city="Rabat")
    assert missing.living_cost_eur is None and missing.monthly_balance_eur is None
    assert any("No sourced living-cost" in n for n in missing.notes)


def test_compensation_confidence_levels():
    assert comp("Pay: $25 per hour.", source="workday:x").confidence == "HIGH"
    assert comp("Pay: $25 per hour.", source="linkedin_search").confidence == "MEDIUM"
    assert comp("Nice office.").confidence == "UNKNOWN"


def test_compensation_score_reweights_missing_factors():
    nominal = comp("Pay: $3,000 per month.")
    assert "nominal_pay" in nominal.score_components and "disposable_balance" not in nominal.score_components
    sourced = comp("Gratification : 6 000 MAD par mois.", country="MA", city="Casablanca")
    assert "disposable_balance" in sourced.score_components and "nominal_pay" not in sourced.score_components
    assert comp("Unpaid internship.").score == 0


# 14-16 Global priority and eligibility ---------------------------------------------------------------
def _gp(c, rejected=False, prefs=None, mode="conversion", rules=None):  # noqa: ANN001
    rules = rules or {"global_priority": yaml.safe_load((PROJECT_ROOT / "config" / "scoring_rules.yaml").read_text())["global_priority"]}
    return priority.compute(80, 85, 50, c, rejected, 70, 40, 60, 45, mode, rules, {"preferences": prefs or {}})


def test_global_priority_weights_and_gates():
    paid = comp("Pay: $3,000 per month.")
    gp = _gp(paid)
    assert gp.score is not None and not gp.provisional and abs(sum(gp.weights.values()) - 1) < 1e-6
    assert gp.weights["recruitability"] == 0.30 and gp.weights["compensation"] == 0.20
    assert _gp(paid, rejected=True).score is None
    assert _gp(comp("Unpaid internship."), prefs={"include_unpaid_internships": False}).filtered
    assert _gp(paid, prefs={"minimum_monthly_salary_eur": 5000}).filtered
    assert _gp(paid, mode="exploration").weights != gp.weights


def test_missing_compensation_gives_a_provisional_score():
    gp = _gp(comp("Nice office."))
    assert gp.provisional and "compensation" not in gp.weights and gp.score is not None
    known = _gp(comp("Pay: $3,000 per month."))
    assert known.score != gp.score


def test_eligibility_filtering_and_unknown_visa(hunter, repo):
    stats = CycleStats()
    raw = RawOpportunity(source="greenhouse:acme", source_id="v1", url="https://x.test/v1", title="Maintenance Engineering Intern",
                         company="Acme", location_text="Singapore",
                         description_text=DATES + "Analyze maintenance data and KPIs. Pay: S$1,800 per month. Python, Excel.")
    opp = hunter.ingest(repo, raw, stats)
    hunter.evaluate(repo, opp, None, stats)
    crit = {c["criterion"]: c["status"] for c in repo.latest_eligibility(opp.id).criteria}
    assert crit["work_authorization"] == "UNKNOWN" and opp.eligibility_status == "NEEDS_VERIFICATION"
    assert opp.global_priority is not None and opp.next_action != "DISCARD"
    strat = repo.latest_fit(opp.id).strategy
    assert strat["caveat"] is None or "work authorization" in strat["caveat"]
    raw2 = raw.model_copy(update={"source_id": "v2", "url": "https://x.test/v2", "location_text": "Paris, France"})
    bad = hunter.ingest(repo, raw2, stats)
    hunter.evaluate(repo, bad, None, stats)
    assert bad.global_priority is None and bad.next_action == "DISCARD"


# 17-19 Semantic, recruitability, networking --------------------------------------------------------------
def test_semantic_classification_from_responsibilities(settings, title_rules):
    text = ("You will run technical demos, support discovery calls with customers, gather customer requirements and "
            "build proof of concept environments with our pre-sales team.")
    fam = title_rules.classify("Customer Solutions Intern")
    m = semantic.classify("Customer Solutions Intern", text, settings.scoring, fam.id)
    assert m.family == "technical_sales" and (m.applied or fam.id == "technical_sales")
    vague = semantic.classify("Summer Intern", "You will run technical demos, discovery calls, proof of concept, rfp answers.",
                              settings.scoring, "unclassified")
    assert vague.applied and vague.family == "technical_sales"
    kept = semantic.classify("Maintenance Engineering Intern", text, settings.scoring, "maintenance_reliability")
    assert not kept.applied                              # a conclusive title is never overridden


def test_recruitability_still_drives_track_b(hunter, repo):
    stats = CycleStats()
    opp = hunter.ingest(repo, RawOpportunity(source="manual", source_id="r1", url="https://x.test/r1", title="Maintenance Engineering Intern",
                                             company="Acme", location_text="Munich, Germany",
                                             description_text=DATES + "MRO team: analyze maintenance data, monitor KPIs, process "
                                             "improvement. Requirements: engineering student, strong data analysis (Excel, Python)."), stats)
    hunter.evaluate(repo, opp, None, stats)
    assert opp.recruitability_score >= 70 and opp.career_track == "industry_robotics"


def test_networking_recommendation_never_invents_a_contact():
    none = networking_reco.recommend("Acme", "Robotics Deployment Intern", [], {}, "industry_robotics")
    assert none["contact"] is None and none["message"] is None
    assert all("linkedin.com/search/results/people" in s["url"] for s in none["searches"])
    contacts = [{"name": "Jane Doe", "relationship": "hiring_manager_post", "confidence": 0.9, "profile_url": "https://linkedin.com/in/jane",
                 "evidence": "post"}, {"name": "Low Conf", "relationship": "recruiter_post", "confidence": 0.3}]
    reco = networking_reco.recommend("Acme", "Robotics Deployment Intern", contacts, {"reasons": ["Air France match."]}, "industry_robotics")
    assert reco["contact"]["name"] == "Jane Doe" and "Hello Jane" in reco["message"] and "Air France" in reco["message"]
    assert "5 business days" in reco["follow_up"]


# 20-22 Signals, deadlines, strategist ---------------------------------------------------------------------
def test_early_signals_are_kept_apart_and_confirmed_later(hunter, repo):
    repo.s.add(LinkedInPost(post_url="https://www.linkedin.com/posts/x-1", text="We are expanding our team, internships coming in 2027",
                            classification="hiring_manager_upcoming_opening", discovery_method="manual"))
    repo.s.add(LinkedInPost(post_url="https://www.linkedin.com/posts/x-2", text="Opening a new plant in Tangier next year",
                            classification="irrelevant", discovery_method="manual"))
    stats = CycleStats()
    pool = hunter.ingest(repo, RawOpportunity(source="manual", source_id="tp", url="https://x.test/tp", title="Internship Talent Community 2027",
                                              company="Zeta", location_text="Casablanca, Morocco", description_text=DATES), stats)
    hunter.evaluate(repo, pool, None, stats)
    signals.sync(repo)
    levels = {s.source_url: s.level for s in signals.recent(repo)}
    assert levels["https://www.linkedin.com/posts/x-1"] == signals.SIGNAL
    assert levels["https://www.linkedin.com/posts/x-2"] == signals.WATCH
    assert levels[pool.official_url] == signals.SIGNAL and pool.next_action == "MONITOR_OPENING"
    real = hunter.ingest(repo, RawOpportunity(source="manual", source_id="real", url="https://x.test/real", title="Maintenance Engineering Intern",
                                              company="Zeta", location_text="Casablanca, Morocco",
                                              description_text=DATES + "maintenance planning and KPIs"), stats)
    hunter.evaluate(repo, real, None, stats)
    signals.sync(repo)
    confirmed = {s.source_url: (s.level, s.opportunity_id) for s in signals.recent(repo)}
    assert confirmed[pool.official_url] == (signals.CONFIRMED, real.id)


def test_opportunity_deadline_reminders_without_repeats(hunter, repo):
    from summer_hunter.crm import followups
    from summer_hunter.notifications.briefing_v3 import opportunity_deadlines

    stats = CycleStats()
    opp = hunter.ingest(repo, RawOpportunity(source="manual", source_id="dl", url="https://x.test/dl", title="Maintenance Engineering Intern",
                                             company="Acme", location_text="Munich, Germany",
                                             description_text=DATES + "Application deadline: 10 November 2026. Maintenance KPIs."), stats)
    hunter.evaluate(repo, opp, None, stats)
    assert opp.application_deadline == date(2026, 11, 10) and opp.deadline_type == "fixed"
    first = opportunity_deadlines(repo, [opp], date(2026, 11, 4))
    assert len(first) == 1 and first[0].key.endswith(":7")
    followups.mark_sent(CRM(repo.s), first)
    assert opportunity_deadlines(repo, [opp], date(2026, 11, 5)) == []
    assert opportunity_deadlines(repo, [opp], date(2026, 11, 11)) == []
    assert opportunity_deadlines(repo, [opp], date(2026, 11, 1)) == []         # posting dates are never deadlines


def test_strategist_recommendations_and_balance(settings):
    rules = settings.scoring
    known = comp("Pay: $3,000 per month.")
    unknown = comp("Nice office.")
    gp_hi = priority.GlobalPriority(score=82, effort_minutes=45)
    gp_mid = priority.GlobalPriority(score=72, provisional=True)
    eligible = [{"criterion": "dates", "status": "PASS", "explanation": "ok"}]
    s1 = strategist.decide(gp_hi, "ELIGIBLE", eligible, known, 85, 88, 20, {"reasons": ["A"]}, {"reasons": ["B"]}, None, None, False, 5, rules)
    assert s1.action == "APPLY_NOW" and s1.category == "SAFETY" and s1.arguments == ["A", "B"]
    s2 = strategist.decide(gp_mid, "ELIGIBLE", eligible, unknown, 85, 70, 20, {}, {}, None, None, False, None, rules)
    assert s2.action == "VERIFY_COMPENSATION"
    s3 = strategist.decide(priority.GlobalPriority(score=None, filtered="failed"), "REJECTED", eligible, known, 1, 1, 1, {}, {}, None,
                           None, False, None, rules)
    assert s3.action == "DISCARD"
    contact = {"name": "Jane", "relationship": "recruiter_post"}
    s4 = strategist.decide(gp_hi, "ELIGIBLE", eligible, known, 80, 70, 95, {}, {}, contact, 1, False, None, rules)
    assert s4.action == "CONTACT_RECRUITER" and s4.category in {"AMBITIOUS", "TARGETED"}
    items = [(c, s) for c, s in [("SAFETY", 90), ("SAFETY", 89), ("SAFETY", 88), ("TARGETED", 70), ("AMBITIOUS", 60)]]
    picked = strategist.balanced(items, 4, {"AMBITIOUS": .25, "TARGETED": .5, "SAFETY": .25}, lambda x: x[0], lambda x: x[1])
    assert {c for c, _ in picked} == {"SAFETY", "TARGETED", "AMBITIOUS"}


# 23 Application studio ----------------------------------------------------------------------------------
def test_application_pack_uses_only_usable_claims(settings):
    from summer_hunter.web.cv import load_content
    from summer_hunter.web.letter import FORBIDDEN
    from summer_hunter.web.pack import talking_points

    claims = knowledge.build(settings.profile_raw, load_content(settings.config_dir))
    statuses = {c.status for c in claims}
    assert {"VERIFIED", "USER_CONFIRMED", "INFERRED", "UNKNOWN"} <= statuses
    usable_texts = {c.text for c in knowledge.usable(claims, "experience")}
    points = talking_points(settings.config_dir, settings.profile_raw, "Maintenance Engineering Intern", "maintenance data KPIs")
    assert len(points) == 3 and all(p.split(": ", 1)[1] in usable_texts for p in points)
    assert len({p.split(": ", 1)[0] for p in points}) == 3                    # three different experiences
    assert "ingénieur d'affaires" in FORBIDDEN


# 24-25 Migrations and Excel ----------------------------------------------------------------------------------
def test_migration_adds_v3_columns_to_an_existing_database(tmp_path):
    import sqlite3

    from summer_hunter.database.migrations import init_db, make_engine

    db = tmp_path / "old.db"
    con = sqlite3.connect(db)
    con.execute("CREATE TABLE opportunities (id INTEGER PRIMARY KEY, canonical_key VARCHAR(64), company_name VARCHAR(300), "
                "normalized_company VARCHAR(200), title VARCHAR(300), normalized_title VARCHAR(300), description_text TEXT, "
                "content_hash VARCHAR(64), official_url VARCHAR(1000), status VARCHAR(20), user_status VARCHAR(20), "
                "is_remote BOOLEAN, notify BOOLEAN, first_discovered_at DATETIME, last_verified_at DATETIME)")
    con.execute("INSERT INTO opportunities (id, canonical_key, company_name, normalized_company, title, normalized_title, "
                "description_text, content_hash, official_url, status, user_status) VALUES (1,'k','A','a','T','t','','h','u','active','new')")
    con.commit()
    con.close()
    init_db(make_engine(f"sqlite:///{db}"))
    cols = {r[1] for r in sqlite3.connect(db).execute("PRAGMA table_info(opportunities)")}
    assert {"compensation_status", "monthly_min_eur", "global_priority", "next_action", "application_deadline"} <= cols
    assert sqlite3.connect(db).execute("SELECT title FROM opportunities").fetchone() == ("T",)      # data kept


def test_excel_has_compensation_and_performance_sheets(hunter, repo, settings, tmp_path):
    from openpyxl import load_workbook

    from summer_hunter.crm.excel_export import FILENAME, export_workbook

    stats = CycleStats()
    opp = hunter.ingest(repo, RawOpportunity(source="greenhouse:acme", source_id="x1", url="https://x.test/x1", title="Maintenance Engineering Intern",
                                             company="Acme", location_text="Munich, Germany",
                                             description_text=DATES + "Maintenance KPIs. Pay: €1,500 per month gross. Relocation assistance."), stats)
    hunter.evaluate(repo, opp, None, stats)
    CRM(repo.s, settings.data_dir).track(opp, "cli")
    wb = load_workbook(export_workbook(CRM(repo.s), tmp_path / FILENAME, "Europe/Paris"))
    assert wb.sheetnames[:5] == ["Dashboard", "Applications", "Networking", "Opportunities", "Activity Log"]
    assert {"Compensation Analysis", "Hunter Performance"} <= set(wb.sheetnames)
    ws = wb["Applications"]
    row = {ws.cell(row=1, column=c).value: ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)}
    assert row["Compensation Status"] == "PAID_CONFIRMED" and row["Monthly Salary EUR"] == 1500
    assert row["Relocation Benefits"] and row["Global Application Priority Score"] is not None and row["Recommended Next Action"]
    assert ws.cell(row=1, column=ws.max_column).value == "Row Version"


# 26 Telegram ----------------------------------------------------------------------------------------------------
def test_v3_telegram_commands_and_buttons(hunter, repo, settings, sessions):
    stats = CycleStats()
    opp = hunter.ingest(repo, RawOpportunity(source="greenhouse:acme", source_id="t1", url="https://x.test/t1", title="Robotics Deployment Intern",
                                             company="Acme", location_text="Casablanca, Morocco",
                                             description_text=DATES + "Deploy autonomous mobile robots, customer onboarding. Gratification : 6 000 MAD par mois."), stats)
    hunter.evaluate(repo, opp, None, stats)
    repo.s.commit()
    settings.telegram_bot_token, settings.telegram_chat_id = TOKEN, CHAT
    fake = FakeTelegram()
    bot = Bot(settings, sessions, TelegramClient(TOKEN, transport=httpx.MockTransport(fake)), NullProvider())
    for cmd in ("/paid", "/highest_paid", "/best_value", "/africa", "/compensation_unknown", "/priority", "/networking",
                "/early_signals", "/hunter_stats", "/mode", "/mode exploration", "/top10", "/crm"):
        bot.handle_update({"update_id": 1, "message": {"chat": {"id": int(CHAT)}, "text": cmd}})
    texts = "\n".join(m["text"] for m in fake.sent)
    for expected in ("Stages payés", "Afrique", "Robotics Deployment Intern", "Priorité globale", "Performance du chasseur",
                     "Mode <b>exploration</b>", "Top 10", "CRM candidatures", "coût de la vie"):
        assert expected in texts, expected
    for data in (f"vc:{opp.id}", f"ve:{opp.id}", f"rt:{opp.id}", f"rr:{opp.id}:EXCELLENT"):
        bot.handle_update({"update_id": 2, "callback_query": {"id": "cb", "data": data, "message": {"message_id": 5, "chat": {"id": int(CHAT)}}}})
    texts = "\n".join(m["text"] for m in fake.sent)
    assert "Rémunération #" in texts and "À vérifier avant de postuler" in texts and "Note la recommandation" in texts
    from summer_hunter.notifications.formatter import buttons as card_buttons

    labels = {b["text"] for row in card_buttons(opp, None) for b in row}
    assert {"Save", "Track", "🧾 Generate application", "💶 Compensation", "✅ Verify eligibility", "🤝 Contact recruiter",
            "Applied", "Dismiss"} <= labels
    assert all(len(b.get("callback_data", "")) <= 64 for row in card_buttons(opp, None) for b in row)


def test_card_shows_compensation_priority_and_recommendation(hunter, repo):
    from summer_hunter.notifications import formatter

    stats = CycleStats()
    opp = hunter.ingest(repo, RawOpportunity(source="greenhouse:acme", source_id="c1", url="https://x.test/c1", title="Robotics Deployment Intern",
                                             company="Acme", location_text="Munich, Germany",
                                             description_text=DATES + "Deploy robots. Pay: €1,200 per month."), stats)
    hunter.evaluate(repo, opp, None, stats)
    fit, ev = repo.latest_fit(opp.id), repo.latest_eligibility(opp.id)
    text, _ = formatter.card(opp, ev.criteria, fit.report, fit.networking, fit.priority, [], None, fit.recruitability,
                             fit.compensation, fit.strategy)
    assert "Compensation:</b> €1,200/month" in text and "Compensation Confidence:</b> HIGH" in text
    assert "Global Priority:" in text and "Recommendation:" in text


# 27-30 Duplicates, CRM status, feedback, metrics ---------------------------------------------------------------
def test_duplicate_postings_merge_into_one_opportunity(hunter, repo):
    stats = CycleStats()
    raw = RawOpportunity(source="greenhouse:acme", source_id="d1", url="https://x.test/d1", title="Maintenance Engineering Intern",
                         company="Acme", location_text="Munich, Germany", description_text=DATES + "x")
    a = hunter.ingest(repo, raw, stats)
    b = hunter.ingest(repo, raw.model_copy(update={"source": "linkedin_search", "source_id": "99", "url": "https://linkedin.com/jobs/view/99/"}), stats)
    assert a.id == b.id and len(a.sources) == 2


def test_application_status_tracking_still_requires_confirmation(hunter, repo, settings):
    stats = CycleStats()
    opp = hunter.ingest(repo, RawOpportunity(source="manual", source_id="s1", url="https://x.test/s1", title="Maintenance Engineering Intern",
                                             company="Acme", location_text="Munich, Germany", description_text=DATES + "x"), stats)
    crm = CRM(repo.s, settings.data_dir)
    app, _ = crm.track(opp, "telegram")
    assert app.status == "TO_APPLY" and app.date_applied is None


def test_rating_learning_needs_enough_evidence(hunter, repo, settings):
    from summer_hunter.engine.feedback import recompute

    stats = CycleStats()
    opps = [hunter.ingest(repo, RawOpportunity(source="manual", source_id=f"f{i}", url=f"https://x.test/f{i}", title="Robotics Deployment Intern",
                                               company=f"Co{i}", location_text="Munich, Germany", description_text=DATES), stats) for i in range(3)]
    for o in opps:
        hunter.evaluate(repo, o, None, stats)
    repo.record_feedback(opps[0], "rating", "WRONG_CAREER_PATH")
    repo.record_feedback(opps[1], "rating", "WRONG_CAREER_PATH")
    assert not [p for p in recompute(repo, settings.scoring) if p["dimension"] == "job_family"]
    repo.record_feedback(opps[2], "rating", "WRONG_CAREER_PATH")
    rows = {(p["dimension"], p["key"]): p["delta"] for p in recompute(repo, settings.scoring)}
    assert rows[("job_family", "robotics_raas")] < 0
    assert repo.latest_eligibility(opps[0].id) is not None                                  # eligibility untouched


def test_hunter_metrics_and_small_samples(hunter, repo):
    stats = CycleStats()
    opp = hunter.ingest(repo, RawOpportunity(source="greenhouse:acme", source_id="m1", url="https://x.test/m1", title="Maintenance Engineering Intern",
                                             company="Acme", location_text="Munich, Germany",
                                             description_text=DATES + "Maintenance KPIs. Pay: €1,500 per month."), stats)
    hunter.evaluate(repo, opp, None, stats)
    repo.record_feedback(opp, "rating", "RELEVANT")
    p = analytics.compute(repo.s, REGIONS)
    assert p.relevant >= 1 and p.paid_confirmed >= 1
    assert p.paid_rate.value is not None and p.paid_rate.small and "indicative" in p.paid_rate.label()
    assert p.precision_at_10.numerator == 1 and p.precision_at_10.denominator == 1


def test_failing_sources_back_off():
    now = datetime(2026, 10, 23, 12, tzinfo=timezone.utc)
    h = SourceHealth(source="x", consecutive_failures=2, last_failure_at=now - timedelta(minutes=5))
    assert not in_backoff(h, now)
    h.consecutive_failures = 4
    assert in_backoff(h, now) and not in_backoff(h, now + timedelta(hours=3))   # 2 h rest, then retried
    h.consecutive_failures = 20
    assert in_backoff(h, now + timedelta(hours=47)) and not in_backoff(h, now + timedelta(hours=49))


def test_mode_switch_changes_weights_through_the_same_evaluation(hunter, repo):
    from summer_hunter.engine import modes

    stats = CycleStats()
    opp = hunter.ingest(repo, RawOpportunity(source="manual", source_id="mo", url="https://x.test/mo", title="Maintenance Engineering Intern",
                                             company="Acme", location_text="Munich, Germany", description_text=DATES + "maintenance KPIs"), stats)
    hunter.evaluate(repo, opp, None, stats)
    before, weights_before = opp.global_priority, repo.latest_fit(opp.id).strategy["global"]["weights"]
    modes.set_mode(repo, "exploration")
    hunter.evaluate(repo, opp, None, stats)               # what /mode triggers for every opportunity, in the background
    after = repo.latest_fit(opp.id).strategy["global"]
    assert after["weights"] != weights_before and opp.global_priority != before
    assert opp.next_action == repo.latest_fit(opp.id).strategy["action"]                 # action recomputed with the score
    with pytest.raises(ValueError):
        modes.set_mode(repo, "turbo")


def test_review_regressions_on_pay_parsing():
    cases = {
        "Visa support is not provided. Pay: $3,000 per month.": ("visa_support", False),
        "No relocation assistance. Pay: $3,000 per month.": ("relocation", False),
    }
    for text, (benefit, present) in cases.items():
        assert (benefit in comp(text).benefits) is present
    unpaid = comp("This unpaid internship provides a $500 relocation stipend.")
    assert unpaid.status == UNPAID_CONFIRMED and unpaid.pay is None
    assert comp("Benefits include a monthly meal allowance. Salary is competitive: $3,000 per month").pay.amount_min == 3000
    assert comp("Salary CA$1,500 per month", country="CA").pay.currency == "CAD"
    assert comp("Pay US$25 per hour", country="CA").pay.currency == "USD"
    assert comp("Salary $3,000 per month", country=None).pay is None              # bare $ without a country: unknown
    assert comp("Salary INR 1,50,000 per month", country="IN").pay.amount_min == 150000
    assert comp("Base salary is $2,000. Meal allowance per month is separate.").pay.period is None
    assert comp("Gross salary details follow. Pay is €1,000 net per month", country="DE").pay.basis == "net"
    ranged = comp("$25/hour, 20-40 hours/week")
    assert (ranged.monthly_min, ranged.monthly_max, ranged.monthly_is_estimate) == (pytest.approx(25 * 20 * 4.333, rel=1e-3),
                                                                                    pytest.approx(25 * 40 * 4.333, rel=1e-3), False)
    assert comp("Overtime is not paid. Salary $3,000 per month.").status == PAID_CONFIRMED
    assert comp("We do not offer unpaid volunteer roles. This internship is unpaid.").status == UNPAID_CONFIRMED
    partial = comp("Pay 6000 MAD par mois", country="MA", city="Casablanca",
                   cfg=CFG | {"living_costs": [{"place": "casablanca, ma", "housing": 400, "source": "x"}]})
    assert partial.living_cost_eur is None                                        # incomplete row: no fake zeros
    bad_scale = comp("Pay: $3,000 per month.", cfg=CFG | {"scoring": CFG["scoring"] | {"monthly_scale_eur": [0, 0]}})
    assert bad_scale.score is not None


def test_floor_uses_the_range_minimum_and_filtered_items_stay_out(hunter, repo):
    gp = _gp(comp("Pay: €1,000 - €3,000 per month.", country="DE"), prefs={"minimum_monthly_salary_eur": 2000})
    assert gp.filtered and gp.score is None
    stats = CycleStats()
    opp = hunter.ingest(repo, RawOpportunity(source="manual", source_id="un", url="https://x.test/un", title="Maintenance Engineering Intern",
                                             company="Acme", location_text="Munich, Germany",
                                             description_text=DATES + "This internship is unpaid. Maintenance KPIs."), stats)
    hunter.evaluate(repo, opp, None, stats)
    assert opp.next_action == "DISCARD" and opp not in repo.top(50) and opp not in repo.top_global(50)
